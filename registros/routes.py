# -*- coding: utf-8 -*-
from __future__ import annotations

import io
import re
from datetime import date
from typing import Any, Dict, List, Sequence, Tuple

from flask import request, jsonify, render_template, send_file, session, abort

from db import conectar_db
from . import registros_bp

from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.pdfbase.pdfmetrics import stringWidth


try:
    from admin.modulos import require_permission
except Exception:
    def require_permission(modulo_codigo: str, acao: str = "ver"):
        def deco(fn):
            return fn
        return deco


try:
    from log import registrar_log, log_erro
except Exception:
    def registrar_log(*args, **kwargs): pass
    def log_erro(*args, **kwargs): pass


# =============================================================================
# CONTEXTO MULTI-CLÍNICA
# =============================================================================

def _clinica_id_atual() -> int:
    clinica_id = session.get("clinica_id")

    if not clinica_id:
        abort(403)

    try:
        return int(clinica_id)
    except Exception:
        abort(403)


def _usuario_id_atual():
    return session.get("usuario_id") or session.get("user_id") or session.get("id")


# =============================================================================
# HELPERS GERAIS · POSTGRES
# =============================================================================

def _safe_str(v) -> str:
    return ("" if v is None else str(v)).strip()


def _only_digits(s: str | None) -> str:
    return re.sub(r"\D+", "", s or "")


def _norm_date_param(s: str) -> str:
    s = _safe_str(s)
    if not s:
        return ""
    if re.match(r"^\d{2}/\d{2}/\d{4}$", s):
        d, m, y = s.split("/")
        return f"{y}-{m}-{d}"
    return s[:10]


def _valid_ident(name: str) -> bool:
    return bool(re.match(r"^[a-zA-Z_][a-zA-Z0-9_]*$", name or ""))


def _row_get(row: Any, key: str, idx: int | None = None, default: Any = None) -> Any:
    if row is None:
        return default
    if isinstance(row, dict):
        return row.get(key, default)
    if idx is not None:
        try:
            return row[idx]
        except Exception:
            return default
    return default


def _rows_to_dicts(cur, rows) -> list[dict]:
    names = [d[0] for d in cur.description] if cur.description else []
    out = []

    for r in rows or []:
        if isinstance(r, dict):
            out.append(dict(r))
        else:
            out.append({names[i]: r[i] for i in range(min(len(names), len(r)))})

    return out


def _get_columns(conn, table: str) -> List[str]:
    if not _valid_ident(table):
        return []

    cur = conn.cursor()
    cur.execute("""
        SELECT column_name
          FROM information_schema.columns
         WHERE table_schema = 'public'
           AND table_name = %s
         ORDER BY ordinal_position
    """, (table,))

    return [_row_get(r, "column_name", 0, "") for r in cur.fetchall() or []]


def _first_existing(cols: Sequence[str], opts: Sequence[str]) -> str | None:
    for c in opts:
        if c in cols:
            return c
    return None


def _has_table(conn, name: str) -> bool:
    if not _valid_ident(name):
        return False

    cur = conn.cursor()
    cur.execute("""
        SELECT EXISTS (
            SELECT 1
              FROM information_schema.tables
             WHERE table_schema = 'public'
               AND table_name = %s
        )
    """, (name,))

    return bool(_row_get(cur.fetchone(), "exists", 0, False))


def _has_column(conn, table: str, col: str) -> bool:
    if not _valid_ident(table) or not _valid_ident(col):
        return False

    cur = conn.cursor()
    cur.execute("""
        SELECT EXISTS (
            SELECT 1
              FROM information_schema.columns
             WHERE table_schema = 'public'
               AND table_name = %s
               AND column_name = %s
        )
    """, (table, col))

    return bool(_row_get(cur.fetchone(), "exists", 0, False))


def _sql_date_expr(col_sql: str) -> str:
    return f"""
    (
      CASE
        WHEN {col_sql} IS NULL THEN NULL
        WHEN {col_sql}::text ~ '^\\d{{4}}-\\d{{2}}-\\d{{2}}' THEN ({col_sql})::date
        WHEN {col_sql}::text ~ '^\\d{{2}}/\\d{{2}}/\\d{{4}}' THEN TO_DATE(SUBSTRING({col_sql}::text FROM 1 FOR 10), 'DD/MM/YYYY')
        ELSE NULL
      END
    )
    """.strip()


def _pick_att_col(cols: list[str] | set[str], *names: str) -> str | None:
    for n in names:
        if n in cols:
            return n
    return None


# =============================================================================
# SCHEMA / ÍNDICES MULTI-CLÍNICA
# =============================================================================

def _ensure_evolucoes_ocultas_schema(conn) -> None:
    cur = conn.cursor()

    cur.execute("""
        CREATE TABLE IF NOT EXISTS atendimento_evolucoes_ocultas (
            id SERIAL PRIMARY KEY,
            clinica_id INTEGER,
            atendimento_id INTEGER NOT NULL REFERENCES atendimentos(id) ON DELETE CASCADE,
            paciente_id INTEGER,
            profissional_id INTEGER,
            profissional_nome TEXT,
            profissional_cbo TEXT,
            evolucao_oculta TEXT NOT NULL,
            visibilidade TEXT NOT NULL DEFAULT 'somente_eu',
            cbos_autorizados TEXT,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            atualizado_em TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
    """)

    cur.execute("""
        ALTER TABLE atendimento_evolucoes_ocultas
        ADD COLUMN IF NOT EXISTS clinica_id INTEGER
    """)

    cur.execute("""
        CREATE INDEX IF NOT EXISTS idx_evo_oculta_clinica
        ON atendimento_evolucoes_ocultas (clinica_id)
    """)

    cur.execute("""
        CREATE INDEX IF NOT EXISTS idx_evo_oculta_clinica_atendimento
        ON atendimento_evolucoes_ocultas (clinica_id, atendimento_id)
    """)

    cur.execute("""
        CREATE INDEX IF NOT EXISTS idx_evo_oculta_atendimento
        ON atendimento_evolucoes_ocultas (atendimento_id)
    """)

    cur.execute("""
        CREATE INDEX IF NOT EXISTS idx_evo_oculta_paciente
        ON atendimento_evolucoes_ocultas (paciente_id)
    """)

    conn.commit()


def ensure_registros_schema(conn) -> None:
    if not _has_table(conn, "atendimentos"):
        return

    cur = conn.cursor()

    cur.execute("""
        ALTER TABLE atendimentos
        ADD COLUMN IF NOT EXISTS clinica_id INTEGER
    """)

    cur.execute("""
        CREATE INDEX IF NOT EXISTS idx_atendimentos_clinica
        ON atendimentos(clinica_id)
    """)

    cols = _get_columns(conn, "atendimentos")

    data_col = _first_existing(cols, ["data_atendimento", "data", "data_iso", "created_at"])
    paciente_col = _first_existing(cols, ["paciente_id"])
    prof_col = _first_existing(cols, ["profissional_id", "prof_id", "id_profissional"])

    if data_col:
        cur.execute(f"""
            CREATE INDEX IF NOT EXISTS idx_atendimentos_clinica_data
            ON atendimentos(clinica_id, {data_col})
        """)

    if paciente_col:
        cur.execute(f"""
            CREATE INDEX IF NOT EXISTS idx_atendimentos_clinica_paciente
            ON atendimentos(clinica_id, {paciente_col})
        """)

    if prof_col:
        cur.execute(f"""
            CREATE INDEX IF NOT EXISTS idx_atendimentos_clinica_profissional
            ON atendimentos(clinica_id, {prof_col})
        """)

    conn.commit()

    _ensure_evolucoes_ocultas_schema(conn)


# =============================================================================
# USUÁRIO LOGADO / EVOLUÇÃO OCULTA
# =============================================================================

def _resolve_logged_user(conn) -> dict:
    uid = None

    for key in ("usuario_id", "user_id", "id"):
        val = session.get(key)
        if val:
            try:
                uid = int(val)
                break
            except Exception:
                pass

    login_like = (
        session.get("usuario_logado")
        or session.get("login")
        or session.get("username")
        or session.get("email")
    )

    if not _has_table(conn, "usuarios"):
        return {"id": uid, "nome": "", "cbo": ""}

    cols = _get_columns(conn, "usuarios")
    cur = conn.cursor()

    nome_expr = "COALESCE(nome, '')" if "nome" in cols else "''"
    cbo_expr = "COALESCE(cbo, '')" if "cbo" in cols else "''"

    if uid:
        cur.execute(f"""
            SELECT id, {nome_expr} AS nome, {cbo_expr} AS cbo
              FROM usuarios
             WHERE id = %s
             LIMIT 1
        """, (uid,))

        r = cur.fetchone()
        if r:
            return {
                "id": _row_get(r, "id", 0),
                "nome": _row_get(r, "nome", 1, "") or "",
                "cbo": _only_digits(_row_get(r, "cbo", 2, "") or ""),
            }

    if login_like:
        busca_cols = [c for c in ("login", "nome", "email") if c in cols]
        if busca_cols:
            conds = [f"LOWER(TRIM(COALESCE({c}, ''))) = LOWER(TRIM(%s))" for c in busca_cols]
            params = [login_like] * len(busca_cols)

            cur.execute(f"""
                SELECT id, {nome_expr} AS nome, {cbo_expr} AS cbo
                  FROM usuarios
                 WHERE {" OR ".join(conds)}
                 LIMIT 1
            """, params)

            r = cur.fetchone()
            if r:
                return {
                    "id": _row_get(r, "id", 0),
                    "nome": _row_get(r, "nome", 1, "") or "",
                    "cbo": _only_digits(_row_get(r, "cbo", 2, "") or ""),
                }

    return {"id": uid, "nome": "", "cbo": ""}


def _hidden_evo_lateral_sql() -> str:
    return """
    LEFT JOIN LATERAL (
        SELECT
            COUNT(*) AS evo_oculta_total,
            COUNT(*) FILTER (
                WHERE
                    (
                        eo.visibilidade = 'somente_eu'
                        AND eo.profissional_id = %s
                    )
                    OR
                    (
                        eo.visibilidade = 'cbos'
                        AND %s <> ''
                        AND POSITION(
                            ',' || %s || ','
                            IN ',' || REGEXP_REPLACE(COALESCE(eo.cbos_autorizados, ''), '\\s+', '', 'g') || ','
                        ) > 0
                    )
            ) AS evo_oculta_visivel,
            STRING_AGG(
                CASE
                    WHEN (
                        eo.visibilidade = 'somente_eu'
                        AND eo.profissional_id = %s
                    )
                    OR
                    (
                        eo.visibilidade = 'cbos'
                        AND %s <> ''
                        AND POSITION(
                            ',' || %s || ','
                            IN ',' || REGEXP_REPLACE(COALESCE(eo.cbos_autorizados, ''), '\\s+', '', 'g') || ','
                        ) > 0
                    )
                    THEN eo.evolucao_oculta
                    ELSE NULL
                END,
                E'\\n---\\n'
                ORDER BY eo.id
            ) AS evolucoes_ocultas_visiveis
        FROM atendimento_evolucoes_ocultas eo
        WHERE eo.atendimento_id = a.id
          AND (
                eo.clinica_id IS NULL
                OR eo.clinica_id = %s
              )
    ) eoh ON TRUE
    """


def _hidden_evo_params(user: dict, clinica_id: int) -> list[Any]:
    uid = user.get("id") or 0
    cbo = _only_digits(user.get("cbo") or "")
    return [uid, cbo, cbo, uid, cbo, cbo, clinica_id]


# =============================================================================
# QUERY COMPARTILHADA
# =============================================================================

def _montar_query_atendimentos(
    conn,
    q: str,
    prof: str,
    data_ini: str,
    data_fim: str,
    status: str,
    sexo: str,
    cid: str,
    cidade: str,
    limit: int | None,
    incluir_evolucao_oculta: bool = True,
) -> Tuple[str, List[Any]]:
    table = "atendimentos"
    cols = _get_columns(conn, table)

    clinica_id = _clinica_id_atual()

    has_proc = _has_table(conn, "atendimento_procedimentos")
    proc_cols = _get_columns(conn, "atendimento_procedimentos") if has_proc else []

    user = _resolve_logged_user(conn)

    nome_cols_raw = [c for c in ["paciente_nome", "nome_paciente", "nome"] if c in cols]
    cpf_cols_raw = [c for c in ["cpf", "paciente_cpf"] if c in cols]
    cns_cols_raw = [c for c in ["cns", "paciente_cns", "cartao_sus"] if c in cols]

    prof_id_col_raw = _first_existing(cols, ["profissional_id", "prof_id", "id_profissional"])
    prof_nome_col_raw = _first_existing(cols, ["profissional_nome", "nome_profissional", "profissional", "usuario_nome", "nome_usuario"])

    data_col_raw = _first_existing(cols, ["data_atendimento", "data", "data_iso", "created_at"])
    status_col_raw = _first_existing(cols, ["status", "situacao", "comparecimento"])
    sexo_col_raw = _first_existing(cols, ["sexo", "sex"])
    cid_col_raw = _first_existing(cols, ["cid", "cid_principal", "cid10"])
    cidade_col_raw = _first_existing(cols, ["cidade", "municipio", "cidade_paciente", "municipio_paciente"])

    nome_cols = [f"a.{c}" for c in nome_cols_raw]
    cpf_cols = [f"a.{c}" for c in cpf_cols_raw]
    cns_cols = [f"a.{c}" for c in cns_cols_raw]

    prof_id_col = f"a.{prof_id_col_raw}" if prof_id_col_raw else None
    prof_nome_col = f"a.{prof_nome_col_raw}" if prof_nome_col_raw else None
    data_col = f"a.{data_col_raw}" if data_col_raw else None
    status_col = f"a.{status_col_raw}" if status_col_raw else None
    sexo_col = f"a.{sexo_col_raw}" if sexo_col_raw else None
    cid_col = f"a.{cid_col_raw}" if cid_col_raw else None
    cidade_col = f"a.{cidade_col_raw}" if cidade_col_raw else None

    joins: List[str] = []
    params: List[Any] = []

    if has_proc:
        joins.append("LEFT JOIN atendimento_procedimentos ap ON ap.atendimento_id = a.id")

    if incluir_evolucao_oculta and _has_table(conn, "atendimento_evolucoes_ocultas"):
        joins.append(_hidden_evo_lateral_sql())
        params.extend(_hidden_evo_params(user, clinica_id))

    where_parts: List[str] = []

    if "clinica_id" in cols:
        where_parts.append("a.clinica_id = %s")
        params.append(clinica_id)

    if q:
        like = f"%{q}%"
        sub: List[str] = []

        for c in nome_cols:
            sub.append(f"{c} ILIKE %s")
            params.append(like)

        for c in cpf_cols:
            sub.append(f"{c}::text ILIKE %s")
            params.append(like)

        for c in cns_cols:
            sub.append(f"{c}::text ILIKE %s")
            params.append(like)

        if has_proc:
            if "procedimento" in proc_cols:
                sub.append("ap.procedimento ILIKE %s")
                params.append(like)
            if "codigo_sigtap" in proc_cols:
                sub.append("ap.codigo_sigtap ILIKE %s")
                params.append(like)

        if incluir_evolucao_oculta and _has_table(conn, "atendimento_evolucoes_ocultas"):
            sub.append("COALESCE(eoh.evolucoes_ocultas_visiveis, '') ILIKE %s")
            params.append(like)

        if sub:
            where_parts.append("(" + " OR ".join(sub) + ")")

    if prof:
        if prof.isdigit() and prof_id_col:
            where_parts.append(f"{prof_id_col} = %s")
            params.append(int(prof))
        elif prof_nome_col:
            where_parts.append(f"{prof_nome_col} ILIKE %s")
            params.append(f"%{prof}%")

    data_ini = _norm_date_param(data_ini)
    data_fim = _norm_date_param(data_fim)

    if data_col and (data_ini or data_fim):
        d_expr = _sql_date_expr(data_col)
        if data_ini and data_fim:
            where_parts.append(f"{d_expr} BETWEEN %s::date AND %s::date")
            params.extend([data_ini, data_fim])
        elif data_ini:
            where_parts.append(f"{d_expr} = %s::date")
            params.append(data_ini)
        else:
            where_parts.append(f"{d_expr} = %s::date")
            params.append(data_fim)

    if status and status_col:
        where_parts.append(f"{status_col} = %s")
        params.append(status)

    if sexo and sexo_col:
        where_parts.append(f"{sexo_col} = %s")
        params.append(sexo)

    if cid and cid_col:
        where_parts.append(f"{cid_col} ILIKE %s")
        params.append(f"%{cid}%")

    if cidade and cidade_col:
        where_parts.append(f"{cidade_col} ILIKE %s")
        params.append(f"%{cidade}%")

    select_parts = [
        "a.*",
        *(["ap.procedimento AS ap_procedimento",
           "COALESCE(ap.codigo_sigtap, '') AS ap_codigo_sigtap"] if has_proc else []),
    ]

    if incluir_evolucao_oculta and _has_table(conn, "atendimento_evolucoes_ocultas"):
        select_parts.extend([
            "COALESCE(eoh.evo_oculta_total, 0) AS evo_oculta_total",
            "COALESCE(eoh.evo_oculta_visivel, 0) AS evo_oculta_visivel",
            """
            CASE
                WHEN COALESCE(eoh.evo_oculta_total, 0) = 0 THEN 'sem_evolucao_oculta'
                WHEN COALESCE(eoh.evo_oculta_visivel, 0) > 0 THEN 'visivel'
                ELSE 'restrita'
            END AS evo_oculta_situacao
            """,
            "COALESCE(eoh.evolucoes_ocultas_visiveis, '') AS evolucoes_ocultas_visiveis",
        ])
    else:
        select_parts.extend([
            "0 AS evo_oculta_total",
            "0 AS evo_oculta_visivel",
            "'sem_evolucao_oculta' AS evo_oculta_situacao",
            "'' AS evolucoes_ocultas_visiveis",
        ])

    sql = f"SELECT {', '.join(select_parts)}\nFROM {table} a"

    if joins:
        sql += "\n" + "\n".join(joins)

    if where_parts:
        sql += "\nWHERE " + " AND ".join(where_parts)

    order_parts: List[str] = []
    if data_col:
        order_parts.append(f"{_sql_date_expr(data_col)} DESC NULLS LAST")
    if "id" in cols:
        order_parts.append("a.id DESC")
    if has_proc:
        order_parts.append("ap.id ASC")

    if order_parts:
        sql += "\nORDER BY " + ", ".join(order_parts)

    if limit and limit > 0:
        sql += "\nLIMIT %s"
        params.append(limit)

    return sql, params


# =============================================================================
# PÁGINA
# =============================================================================

@registros_bp.get("/")
@require_permission("registros", "ver")
def pagina_registros():
    registrar_log(
        modulo="registros",
        acao="visualizar",
        entidade="registros",
        descricao="Abriu tela de registros.",
        detalhes={"clinica_id": session.get("clinica_id")},
    )

    return render_template("registros.html")


# =============================================================================
# API LISTAR
# =============================================================================

@registros_bp.get("/api/list")
@require_permission("registros", "ver")
def api_listar_atendimentos():
    clinica_id = _clinica_id_atual()

    q = (request.args.get("q") or "").strip()
    prof = (request.args.get("prof") or "").strip()
    data_ini = (request.args.get("data_ini") or "").strip()
    data_fim = (request.args.get("data_fim") or "").strip()
    status = (request.args.get("status") or "").strip()
    sexo = (request.args.get("sexo") or "").strip()
    cid = (request.args.get("cid") or "").strip()
    cidade = (request.args.get("cidade") or "").strip()

    try:
        limit = int(request.args.get("limit", 500))
    except ValueError:
        limit = 500

    conn = conectar_db()

    try:
        try:
            conn.rollback()
        except Exception:
            pass

        ensure_registros_schema(conn)

        if not _has_table(conn, "atendimentos"):
            return jsonify([])

        cur = conn.cursor()

        sql, params = _montar_query_atendimentos(
            conn,
            q=q,
            prof=prof,
            data_ini=data_ini,
            data_fim=data_fim,
            status=status,
            sexo=sexo,
            cid=cid,
            cidade=cidade,
            limit=limit,
        )

        cur.execute(sql, params)
        data = _rows_to_dicts(cur, cur.fetchall())

        pac_cache: dict = {}
        ag_cache: dict = {}

        for r in data:
            _enrich_with_paciente(conn, r, pac_cache)
            _enrich_with_agendamento(conn, r, ag_cache)

            r["pac__nome"] = r.get("pac__nome") or r.get("paciente_nome") or r.get("nome") or ""
            r["pac__cpf"] = r.get("pac__cpf") or r.get("paciente_cpf") or r.get("cpf") or ""
            r["pac__cns"] = r.get("pac__cns") or r.get("paciente_cns") or r.get("cns") or r.get("cartao_sus") or ""
            r["pac__nascimento"] = r.get("pac__nascimento") or r.get("paciente_nascimento") or r.get("nascimento") or r.get("data_nascimento") or ""
            r["pac__status"] = r.get("pac__status") or r.get("status") or ""
            r["pac__sexo"] = r.get("pac__sexo") or r.get("sexo") or ""
            r["pac__cid"] = r.get("pac__cid") or r.get("cid") or ""
            r["pac__municipio"] = r.get("pac__municipio") or r.get("cidade") or r.get("municipio") or ""

        registrar_log(
            modulo="registros",
            acao="visualizar",
            entidade="atendimentos",
            descricao="Consultou registros.",
            detalhes={
                "clinica_id": clinica_id,
                "total": len(data),
                "filtros": {
                    "q": q,
                    "prof": prof,
                    "data_ini": data_ini,
                    "data_fim": data_fim,
                    "status": status,
                    "sexo": sexo,
                    "cid": cid,
                    "cidade": cidade,
                    "limit": limit,
                },
            },
        )

        return jsonify(data)

    except Exception as e:
        log_erro(
            "registros",
            e,
            entidade="atendimentos",
            descricao="Erro ao consultar registros.",
            detalhes={"clinica_id": clinica_id},
        )
        return jsonify({"ok": False, "erro": str(e)}), 500

    finally:
        conn.close()


# =============================================================================
# ENRIQUECIMENTOS
# =============================================================================

def _enrich_with_paciente(conn, base_row: dict, cache: dict) -> None:
    if not _has_table(conn, "pacientes"):
        return

    clinica_id = _clinica_id_atual()

    pcols = _get_columns(conn, "pacientes")
    col_id = _pick_att_col(pcols, "id", "paciente_id")
    col_clinica = _pick_att_col(pcols, "clinica_id")
    col_nome = _pick_att_col(pcols, "nome", "paciente_nome", "nome_paciente")
    col_cpf = _pick_att_col(pcols, "cpf", "paciente_cpf", "cpf_digits", "cpf_cidadao")
    col_cns = _pick_att_col(pcols, "cns", "paciente_cns", "cns_digits", "cartao_sus")
    col_nasc = _pick_att_col(pcols, "nascimento", "data_nascimento", "dt_nasc", "paciente_nascimento")
    col_pront = _pick_att_col(pcols, "prontuario", "prontuario_num")
    col_idade = _pick_att_col(pcols, "idade")

    extras = {
        "sexo": _pick_att_col(pcols, "sexo", "sex"),
        "telefone": _pick_att_col(pcols, "telefone", "telefone1", "paciente_telefone1", "celular"),
        "cep": _pick_att_col(pcols, "cep", "paciente_cep"),
        "logradouro": _pick_att_col(pcols, "logradouro", "rua", "paciente_logradouro"),
        "numero": _pick_att_col(pcols, "numero", "numero_casa", "paciente_numero_casa"),
        "bairro": _pick_att_col(pcols, "bairro", "paciente_bairro"),
        "municipio": _pick_att_col(pcols, "municipio", "cidade", "paciente_municipio", "paciente_cidade"),
        "uf": _pick_att_col(pcols, "uf"),
        "mod": _pick_att_col(pcols, "mod", "modalidade"),
        "status": _pick_att_col(pcols, "status", "situacao"),
        "cid": _pick_att_col(pcols, "cid", "cid_principal", "cid10"),
        "cid2": _pick_att_col(pcols, "cid2"),
        "mae": _pick_att_col(pcols, "mae", "nome_mae"),
        "pai": _pick_att_col(pcols, "pai", "nome_pai"),
        "responsavel": _pick_att_col(pcols, "responsavel"),
        "alergias": _pick_att_col(pcols, "alergias"),
        "aviso": _pick_att_col(pcols, "aviso"),
        "comorbidades_json": _pick_att_col(pcols, "comorbidades_json"),
        "raca": _pick_att_col(pcols, "raca"),
        "estado_civil": _pick_att_col(pcols, "estado_civil"),
        "complemento": _pick_att_col(pcols, "complemento"),
    }

    if not col_id:
        return

    pid = base_row.get("paciente_id") or base_row.get("pacienteId") or ""
    cpf = base_row.get("paciente_cpf") or base_row.get("cpf") or ""
    cns = base_row.get("paciente_cns") or base_row.get("cns") or base_row.get("cartao_sus") or ""
    nome = base_row.get("paciente_nome") or base_row.get("nome_paciente") or base_row.get("nome") or ""
    nasc = base_row.get("paciente_nascimento") or base_row.get("nascimento") or base_row.get("data_nascimento") or ""

    cpf_d = _only_digits(str(cpf))
    cns_d = _only_digits(str(cns))
    nasc_iso = _norm_date_param(str(nasc))

    if str(pid).strip():
        cache_key = f"{clinica_id}:pid:{pid}"
    elif cpf_d:
        cache_key = f"{clinica_id}:cpf:{cpf_d}"
    elif cns_d:
        cache_key = f"{clinica_id}:cns:{cns_d}"
    elif str(nome).strip() and nasc_iso:
        cache_key = f"{clinica_id}:nn:{str(nome).strip().lower()}|{nasc_iso}"
    else:
        return

    if cache_key in cache:
        pac = cache[cache_key]
    else:
        cur = conn.cursor()
        fields = [col_id]

        for c in [col_clinica, col_nome, col_cpf, col_cns, col_nasc, col_pront, col_idade, *extras.values()]:
            if c and c not in fields:
                fields.append(c)

        pac = None
        filtro_clinica = f" AND {col_clinica} = %s" if col_clinica else ""

        if str(pid).strip() and str(pid).strip().isdigit():
            params = [int(pid)]
            if col_clinica:
                params.append(clinica_id)

            cur.execute(
                f"SELECT {', '.join(fields)} FROM pacientes WHERE {col_id} = %s {filtro_clinica} LIMIT 1",
                params,
            )
            pac = cur.fetchone()

        if pac is None and col_cpf and cpf_d:
            params = [cpf_d]
            if col_clinica:
                params.append(clinica_id)

            cur.execute(
                f"""
                SELECT {', '.join(fields)}
                  FROM pacientes
                 WHERE REGEXP_REPLACE(COALESCE({col_cpf}::text, ''), '\\D', '', 'g') = %s
                 {filtro_clinica}
                 LIMIT 1
                """,
                params,
            )
            pac = cur.fetchone()

        if pac is None and col_cns and cns_d:
            params = [cns_d]
            if col_clinica:
                params.append(clinica_id)

            cur.execute(
                f"""
                SELECT {', '.join(fields)}
                  FROM pacientes
                 WHERE REGEXP_REPLACE(COALESCE({col_cns}::text, ''), '\\D', '', 'g') = %s
                 {filtro_clinica}
                 LIMIT 1
                """,
                params,
            )
            pac = cur.fetchone()

        if pac is None and col_nome and col_nasc and str(nome).strip() and nasc_iso:
            params = [str(nome).strip(), nasc_iso]
            if col_clinica:
                params.append(clinica_id)

            cur.execute(
                f"""
                SELECT {', '.join(fields)}
                  FROM pacientes
                 WHERE TRIM(LOWER({col_nome})) = TRIM(LOWER(%s))
                   AND {_sql_date_expr(col_nasc)} = %s::date
                 {filtro_clinica}
                 LIMIT 1
                """,
                params,
            )
            pac = cur.fetchone()

        pac_dict = {}
        if pac is not None:
            for i, fname in enumerate(fields):
                pac_dict[fname] = _row_get(pac, fname, i, "")

        cache[cache_key] = pac_dict
        pac = pac_dict

    if not isinstance(pac, dict) or not pac:
        return

    def put(out_key: str, col: str | None):
        if col and col in pac:
            base_row[out_key] = pac.get(col, "")

    put("pac__nome", col_nome)
    put("pac__cpf", col_cpf)
    put("pac__cns", col_cns)
    put("pac__nascimento", col_nasc)
    put("pac__prontuario", col_pront)
    put("pac__idade", col_idade)

    for key, col in extras.items():
        put(f"pac__{key}", col)


def _enrich_with_agendamento(conn, base_row: dict, cache: dict) -> None:
    if not _has_table(conn, "agendamentos"):
        return

    clinica_id = _clinica_id_atual()

    acols = _get_columns(conn, "agendamentos")
    col_clinica = _pick_att_col(acols, "clinica_id")
    col_pid = _pick_att_col(acols, "paciente_id", "cidadao_id")
    col_pnome = _pick_att_col(acols, "paciente", "paciente_nome", "nome_paciente", "nome")
    col_ini = _pick_att_col(acols, "inicio", "datahora", "data_hora", "data_inicio")
    col_prof = _pick_att_col(acols, "profissional", "profissional_nome", "nome_profissional")
    col_profid = _pick_att_col(acols, "profissional_id", "prof_id")
    col_profcbo = _pick_att_col(acols, "profissional_cbo", "cbo")

    if not col_ini:
        return

    dt_raw = base_row.get("data_atendimento") or base_row.get("data") or base_row.get("data_iso") or base_row.get("created_at") or ""
    dt_iso = _norm_date_param(str(dt_raw)[:10])
    if not dt_iso:
        return

    pid = base_row.get("paciente_id") or ""
    nome = base_row.get("paciente_nome") or base_row.get("nome_paciente") or base_row.get("nome") or ""

    if str(pid).strip():
        key = f"{clinica_id}:ag:{pid}|{dt_iso}"
    elif str(nome).strip():
        key = f"{clinica_id}:ag:n:{str(nome).strip().lower()}|{dt_iso}"
    else:
        return

    if key in cache:
        ag = cache[key]
    else:
        cur = conn.cursor()
        ag = None
        filtro_clinica = f" AND {col_clinica} = %s" if col_clinica else ""

        if col_pid and str(pid).strip() and str(pid).strip().isdigit():
            params = [dt_iso, int(pid)]
            if col_clinica:
                params.append(clinica_id)

            cur.execute(
                f"""
                SELECT *
                  FROM agendamentos
                 WHERE {_sql_date_expr(col_ini)} = %s::date
                   AND {col_pid} = %s
                   {filtro_clinica}
                 ORDER BY {col_ini} ASC
                 LIMIT 1
                """,
                params,
            )
            ag = cur.fetchone()

        if ag is None and col_pnome and str(nome).strip():
            params = [dt_iso, str(nome).strip()]
            if col_clinica:
                params.append(clinica_id)

            cur.execute(
                f"""
                SELECT *
                  FROM agendamentos
                 WHERE {_sql_date_expr(col_ini)} = %s::date
                   AND TRIM(LOWER({col_pnome})) = TRIM(LOWER(%s))
                   {filtro_clinica}
                 ORDER BY {col_ini} ASC
                 LIMIT 1
                """,
                params,
            )
            ag = cur.fetchone()

        ag_dict = {}
        if ag is not None:
            names = [d[0] for d in cur.description]
            if isinstance(ag, dict):
                ag_dict = dict(ag)
            else:
                ag_dict = {names[i]: ag[i] for i in range(min(len(names), len(ag)))}

        cache[key] = ag_dict
        ag = ag_dict

    if not isinstance(ag, dict) or not ag:
        return

    ini = ag.get(col_ini, "")
    base_row["ag__inicio"] = ini
    base_row["ag__hora"] = str(ini)[11:16] if isinstance(ini, str) and len(ini) >= 16 else ""

    if col_prof and col_prof in ag:
        base_row["ag__profissional"] = ag.get(col_prof, "")
    if col_profid and col_profid in ag:
        base_row["ag__profissional_id"] = ag.get(col_profid, "")
    if col_profcbo and col_profcbo in ag:
        base_row["ag__prof_cbo"] = ag.get(col_profcbo, "")


# =============================================================================
# XLSX COMPLETO
# =============================================================================

def _normalize_keys_for_xlsx(rows: list[dict]) -> list[str]:
    if not rows:
        return ["Mensagem"]

    preferred = [
        "id", "clinica_id",
        "paciente_id", "paciente_nome", "nome", "nome_paciente",
        "paciente_cpf", "cpf", "pac__cpf",
        "paciente_cns", "cns", "cartao_sus", "pac__cns",
        "paciente_nascimento", "nascimento", "data_nascimento", "pac__nascimento",
        "data_atendimento", "data", "data_iso", "created_at",
        "profissional_id", "prof_id", "id_profissional",
        "profissional_nome", "nome_profissional", "profissional",
        "profissional_cbo", "cbo_profissional", "cbo",
        "status", "situacao", "comparecimento",
        "cid", "cid_principal", "cid10", "pac__cid",
        "cidade", "municipio", "pac__municipio",
        "procedimento", "ap_procedimento",
        "codigo_sigtap", "ap_codigo_sigtap",
        "evolucao",
        "evo_oculta_situacao",
        "evo_oculta_total",
        "evo_oculta_visivel",
        "evolucoes_ocultas_visiveis",
        "pac__prontuario", "pac__idade",
        "pac__telefone", "pac__sexo", "pac__logradouro", "pac__numero", "pac__bairro", "pac__cep",
        "ag__inicio", "ag__hora", "ag__profissional", "ag__profissional_id", "ag__prof_cbo",
    ]

    keys_all: set[str] = set()

    for r in rows:
        keys_all.update(r.keys())

    out: list[str] = []

    for k in preferred:
        if k in keys_all and k not in out:
            out.append(k)

    out.extend(sorted([k for k in keys_all if k not in out]))

    return out


@registros_bp.get("/exportar_xlsx")
@require_permission("registros", "exportar")
def exportar_xlsx():
    clinica_id = _clinica_id_atual()

    filtros = {
        "q": (request.args.get("q") or "").strip(),
        "prof": (request.args.get("prof") or "").strip(),
        "data_ini": (request.args.get("data_ini") or "").strip(),
        "data_fim": (request.args.get("data_fim") or "").strip(),
        "status": (request.args.get("status") or "").strip(),
        "sexo": (request.args.get("sexo") or "").strip(),
        "cid": (request.args.get("cid") or "").strip(),
        "cidade": (request.args.get("cidade") or "").strip(),
    }

    try:
        limit = int(request.args.get("limit", 5000))
    except ValueError:
        limit = 5000

    conn = conectar_db()

    try:
        try:
            conn.rollback()
        except Exception:
            pass

        ensure_registros_schema(conn)

        cur = conn.cursor()
        sql, params = _montar_query_atendimentos(conn, limit=limit, **filtros)
        cur.execute(sql, params)

        data = _rows_to_dicts(cur, cur.fetchall())

        pac_cache: dict = {}
        ag_cache: dict = {}

        for r in data:
            _enrich_with_paciente(conn, r, pac_cache)
            _enrich_with_agendamento(conn, r, ag_cache)

    except Exception as e:
        log_erro(
            "registros",
            e,
            entidade="xlsx",
            descricao="Erro ao exportar registros XLSX.",
            detalhes={"clinica_id": clinica_id, "filtros": filtros},
        )
        return jsonify({"ok": False, "erro": str(e)}), 500

    finally:
        conn.close()

    try:
        from openpyxl import Workbook
    except ImportError:
        return jsonify({"error": "Instale o pacote 'openpyxl' para exportar XLSX."}), 500

    wb = Workbook()
    ws = wb.active
    ws.title = "Registros"

    if not data:
        ws.append(["Mensagem"])
        ws.append(["Nenhum registro para os filtros selecionados."])
    else:
        cols = _normalize_keys_for_xlsx(data)
        ws.append(cols)

        for row in data:
            ws.append([row.get(c, "") for c in cols])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    registrar_log(
        modulo="registros",
        acao="exportar",
        entidade="xlsx",
        descricao="Exportou registros XLSX.",
        detalhes={
            "clinica_id": clinica_id,
            "total": len(data),
            "filtros": filtros,
        },
    )

    filename = f"registros_{date.today().isoformat()}.xlsx"

    return send_file(
        buf,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

# =============================================================================
# PDF EVOLUÇÕES · COM TIMBRE POR CLÍNICA
# =============================================================================

import io
from datetime import date, datetime
from typing import Any, Dict, Optional

from flask import request, jsonify, send_file
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas
from reportlab.pdfbase.pdfmetrics import stringWidth




PDF_FONT = "Helvetica"
PDF_FONT_BOLD = "Helvetica-Bold"


def row_to_dict(row, cursor=None):
    """
    Converte retorno do cursor em dict.
    Funciona com RealDictCursor, sqlite Row e tupla comum.
    """
    if not row:
        return None

    if isinstance(row, dict):
        return dict(row)

    if hasattr(row, "keys"):
        return dict(row)

    if cursor and getattr(cursor, "description", None):
        colunas = [desc[0] for desc in cursor.description]
        return dict(zip(colunas, row))

    return None



def _fmt_cpf(v: str) -> str:
    d = _only_digits(v)
    if len(d) == 11:
        return f"{d[:3]}.{d[3:6]}.{d[6:9]}-{d[9:]}"
    return _safe_str(v) or "—"


def _fmt_cns(v: str) -> str:
    d = _only_digits(v)
    return d if d else (_safe_str(v) or "—")


def _fmt_data_br(v: Any) -> str:
    s = _safe_str(v)
    if not s:
        return "—"

    try:
        if len(s) >= 10 and s[4] == "-" and s[7] == "-":
            return datetime.strptime(s[:10], "%Y-%m-%d").strftime("%d/%m/%Y")
    except Exception:
        pass

    return s[:10]


def _calc_idade_from_iso(nasc_iso: str) -> str:
    s = _safe_str(nasc_iso)
    if not s or len(s) < 10 or s[4] != "-" or s[7] != "-":
        return "—"

    try:
        y, m, d = int(s[:4]), int(s[5:7]), int(s[8:10])
        hoje = date.today()
        idade = hoje.year - y - ((hoje.month, hoje.day) < (m, d))
        return str(max(0, idade))
    except Exception:
        return "—"


def _img_bytes(v):
    if not v:
        return None
    if isinstance(v, memoryview):
        return v.tobytes()
    if isinstance(v, bytearray):
        return bytes(v)
    if isinstance(v, bytes):
        return v
    return None


def _pdf_int(v, default: int) -> int:
    try:
        return int(v if v is not None else default)
    except Exception:
        return default


def _pdf_bool(v) -> bool:
    return v in (True, "true", "1", 1, "on", "sim", "yes")


def _pdf_safe_hex(v: str, default="#0f766e"):
    s = (_safe_str(v) or default).strip()
    if not s.startswith("#"):
        return colors.HexColor(default)
    try:
        return colors.HexColor(s)
    except Exception:
        return colors.HexColor(default)


def _wrap_text(text: str, font_name: str, font_size: int, max_width: float) -> list[str]:
    t = _safe_str(text)
    if not t:
        return ["—"]

    t = t.replace("\r", "")
    paragraphs = [p.strip() for p in t.split("\n")]
    lines: list[str] = []

    for p in paragraphs:
        if p == "":
            lines.append("")
            continue

        words = p.split()
        cur = ""

        for w in words:
            cand = (cur + " " + w).strip() if cur else w
            if stringWidth(cand, font_name, font_size) <= max_width:
                cur = cand
            else:
                if cur:
                    lines.append(cur)
                    cur = w
                else:
                    lines.append(w)

        if cur:
            lines.append(cur)

    return lines or ["—"]


def _cbo_desc(conn, cbo: str) -> str:
    code = _safe_str(cbo)
    if not code:
        return "—"

    for tb in ("cbo_catalogo", "cbos", "ocupacoes", "cbo_funcao", "cbo_descricao"):
        if not _has_table(conn, tb):
            continue

        cols = _get_columns(conn, tb)
        c_codigo = _pick_att_col(cols, "co_ocupacao", "codigo", "cbo", "cod")
        c_desc = _pick_att_col(cols, "no_ocupacao", "descricao", "desc", "funcao", "nome")

        if not (c_codigo and c_desc):
            continue

        cur = conn.cursor()
        cur.execute(
            f"""
            SELECT {c_desc}
              FROM {tb}
             WHERE REGEXP_REPLACE(COALESCE({c_codigo}::text, ''), '\\D', '', 'g') =
                   REGEXP_REPLACE(%s, '\\D', '', 'g')
             LIMIT 1
            """,
            (code,),
        )
        r = cur.fetchone()

        if r and _row_get(r, c_desc, 0, ""):
            return str(_row_get(r, c_desc, 0, "")).strip()

    return "—"


def _buscar_timbre_pdf(conn, clinica_id: int) -> dict:
    """
    Busca o timbre direto do banco para uso no ReportLab.
    Não usa URL; usa binário das imagens.
    """
    cfg = {}

    try:
        cur = conn.cursor()
        cur.execute("""
            SELECT *
              FROM clinica_configuracoes
             WHERE clinica_id = %s
             ORDER BY id DESC
             LIMIT 1
        """, (clinica_id,))
        row = cur.fetchone()
        cfg = row_to_dict(row, cur) or {}
    except Exception:
        cfg = {}

    return {
        "clinica_id": clinica_id,

        "cabecalho_texto": _safe_str(cfg.get("cabecalho_texto")),
        "cabecalho_altura": _pdf_int(cfg.get("cabecalho_altura"), 130),
        "cabecalho_mostrar_logo": _pdf_bool(cfg.get("cabecalho_mostrar_logo")),
        "cabecalho_alinhamento": _safe_str(cfg.get("cabecalho_alinhamento")) or "centro",

        "rodape_texto": _safe_str(cfg.get("rodape_texto")),
        "rodape_altura": _pdf_int(cfg.get("rodape_altura"), 90),
        "rodape_alinhamento": _safe_str(cfg.get("rodape_alinhamento")) or "esquerda",

        "margem_superior": _pdf_int(cfg.get("margem_superior"), 16),
        "margem_inferior": _pdf_int(cfg.get("margem_inferior"), 16),
        "margem_esquerda": _pdf_int(cfg.get("margem_esquerda"), 18),
        "margem_direita": _pdf_int(cfg.get("margem_direita"), 18),

        "mostrar_linha_cabecalho": _pdf_bool(cfg.get("mostrar_linha_cabecalho")),
        "mostrar_linha_rodape": _pdf_bool(cfg.get("mostrar_linha_rodape")),
        "cor_listra_topo": _safe_str(cfg.get("cor_listra_topo")) or "#0f766e",

        "logo_bin": _img_bytes(cfg.get("logo_bin")),
        "cabecalho_img_bin": _img_bytes(cfg.get("cabecalho_img_bin")),
        "rodape_img_bin": _img_bytes(cfg.get("rodape_img_bin")),
        "rodape_img_2_bin": _img_bytes(cfg.get("rodape_img_2_bin")),
        "rodape_img_3_bin": _img_bytes(cfg.get("rodape_img_3_bin")),
    }


def _draw_img_fit(c: canvas.Canvas, img_bytes: bytes, x: float, y: float, w: float, h: float):
    try:
        img = ImageReader(io.BytesIO(img_bytes))
        iw, ih = img.getSize()

        ratio = min(w / iw, h / ih)
        dw = iw * ratio
        dh = ih * ratio

        c.drawImage(
            img,
            x + (w - dw) / 2,
            y + (h - dh) / 2,
            width=dw,
            height=dh,
            preserveAspectRatio=True,
            mask="auto",
        )
    except Exception:
        pass


def _pdf_content_box(timbre: dict, page_w: float, page_h: float):
    mx = timbre["margem_esquerda"] * mm
    mr = timbre["margem_direita"] * mm

    header_h = max(22 * mm, timbre["cabecalho_altura"] * 0.55)
    footer_h = max(18 * mm, timbre["rodape_altura"] * 0.45)

    top_y = page_h - header_h - (timbre["margem_superior"] * mm)
    bottom_y = footer_h + (timbre["margem_inferior"] * mm)

    return {
        "x": mx,
        "w": page_w - mx - mr,
        "top": top_y,
        "bottom": bottom_y,
        "header_h": header_h,
        "footer_h": footer_h,
    }


def _draw_timbre_page(c: canvas.Canvas, timbre: dict, page_w: float, page_h: float):
    stripe = _pdf_safe_hex(timbre.get("cor_listra_topo"))

    c.setFillColor(stripe)
    c.rect(0, page_h - 5 * mm, page_w, 5 * mm, fill=1, stroke=0)

    header_h = max(22 * mm, timbre["cabecalho_altura"] * 0.55)
    footer_h = max(18 * mm, timbre["rodape_altura"] * 0.45)

    header_y = page_h - header_h - 7 * mm

    if timbre.get("cabecalho_img_bin"):
        _draw_img_fit(c, timbre["cabecalho_img_bin"], 12 * mm, header_y, page_w - 24 * mm, header_h)

    elif timbre.get("cabecalho_texto"):
        c.setFont(PDF_FONT_BOLD, 12)
        c.setFillColor(colors.HexColor("#0F172A"))
        lines = _wrap_text(timbre["cabecalho_texto"], PDF_FONT_BOLD, 12, page_w - 40 * mm)[:3]
        y = page_h - 17 * mm
        for ln in lines:
            c.drawCentredString(page_w / 2, y, ln)
            y -= 5 * mm

    if timbre.get("cabecalho_mostrar_logo") and timbre.get("logo_bin"):
        _draw_img_fit(c, timbre["logo_bin"], 15 * mm, page_h - 32 * mm, 28 * mm, 22 * mm)

    if timbre.get("mostrar_linha_cabecalho"):
        c.setStrokeColor(colors.HexColor("#CBD5E1"))
        c.setLineWidth(0.8)
        c.line(15 * mm, header_y - 3 * mm, page_w - 15 * mm, header_y - 3 * mm)

    footer_y = 8 * mm

    rodapes = [
        timbre.get("rodape_img_bin"),
        timbre.get("rodape_img_2_bin"),
        timbre.get("rodape_img_3_bin"),
    ]
    rodapes = [r for r in rodapes if r]

    if rodapes:
        slot_w = (page_w - 24 * mm) / len(rodapes)
        for i, img in enumerate(rodapes):
            _draw_img_fit(c, img, 12 * mm + i * slot_w, footer_y, slot_w, footer_h)
    elif timbre.get("rodape_texto"):
        c.setFont(PDF_FONT, 8)
        c.setFillColor(colors.HexColor("#475569"))
        lines = _wrap_text(timbre["rodape_texto"], PDF_FONT, 8, page_w - 36 * mm)[:3]
        y = footer_y + footer_h - 5 * mm
        for ln in lines:
            if timbre.get("rodape_alinhamento") == "centro":
                c.drawCentredString(page_w / 2, y, ln)
            else:
                c.drawString(18 * mm, y, ln)
            y -= 4 * mm

    if timbre.get("mostrar_linha_rodape"):
        c.setStrokeColor(colors.HexColor("#CBD5E1"))
        c.setLineWidth(0.8)
        c.line(15 * mm, footer_y + footer_h + 2 * mm, page_w - 15 * mm, footer_y + footer_h + 2 * mm)


def _draw_doc_title(c: canvas.Canvas, x: float, y: float, w: float, titulo: str, subtitulo: str):
    c.setFillColor(colors.HexColor("#0F172A"))
    c.setFont(PDF_FONT_BOLD, 16)
    c.drawString(x, y, titulo)

    c.setFillColor(colors.HexColor("#64748B"))
    c.setFont(PDF_FONT, 9)
    c.drawRightString(x + w, y + 1 * mm, subtitulo)

    c.setStrokeColor(colors.HexColor("#E2E8F0"))
    c.setLineWidth(1)
    c.line(x, y - 4 * mm, x + w, y - 4 * mm)

    return y - 11 * mm


def _draw_patient_box(c: canvas.Canvas, x: float, y: float, w: float, pac: dict, tipo_pdf: str):
    h = 24 * mm

    c.setFillColor(colors.HexColor("#F8FAFC"))
    c.setStrokeColor(colors.HexColor("#E2E8F0"))
    c.roundRect(x, y - h, w, h, 8, fill=1, stroke=1)

    c.setFillColor(colors.HexColor("#0F172A"))
    c.setFont(PDF_FONT_BOLD, 12)
    c.drawString(x + 7 * mm, y - 8 * mm, (_safe_str(pac.get("nome")) or "—")[:86])

    idade = _calc_idade_from_iso(pac.get("nasc", ""))

    c.setFont(PDF_FONT, 8.7)
    c.setFillColor(colors.HexColor("#334155"))

    linha1 = f"Idade: {idade} anos   •   CPF: {pac.get('cpf', '—')}   •   CNS: {pac.get('cns', '—')}"
    linha2 = f"Modalidade: {pac.get('mod', '—')}   •   Status: {pac.get('status', '—')}   •   Relatório: {tipo_pdf}"

    c.drawString(x + 7 * mm, y - 14 * mm, linha1[:130])
    c.drawString(x + 7 * mm, y - 19 * mm, linha2[:130])

    return y - h - 7 * mm


def _registro_height(evolucao: str, w: float) -> float:
    inner_w = w - 14 * mm
    lines = _wrap_text(evolucao, PDF_FONT, 9, inner_w)
    qtd = min(len(lines), 13)
    return max(34 * mm, (22 * mm) + qtd * 4.2 * mm)


def _draw_registro_pdf(
    c: canvas.Canvas,
    x: float,
    y: float,
    w: float,
    *,
    profissional: str,
    cbo: str,
    cbo_desc: str,
    data_atendimento: str,
    evolucao: str,
    badge: str = "",
):
    h = _registro_height(evolucao, w)
    pad = 7 * mm
    inner_w = w - 2 * pad

    c.setFillColor(colors.white)
    c.setStrokeColor(colors.HexColor("#E5E7EB"))
    c.roundRect(x, y - h, w, h, 8, fill=1, stroke=1)

    c.setFont(PDF_FONT_BOLD, 9.5)
    c.setFillColor(colors.HexColor("#0F172A"))
    c.drawString(x + pad, y - 7 * mm, (_safe_str(profissional) or "—")[:70])

    c.setFont(PDF_FONT, 8)
    c.setFillColor(colors.HexColor("#64748B"))
    c.drawString(x + pad, y - 12 * mm, f"CBO: {_safe_str(cbo) or '—'} • {_safe_str(cbo_desc) or '—'}"[:105])

    c.setFont(PDF_FONT_BOLD, 8.5)
    c.setFillColor(colors.HexColor("#0F766E"))
    c.drawRightString(x + w - pad, y - 7 * mm, _fmt_data_br(data_atendimento))

    if badge:
        c.setFont(PDF_FONT_BOLD, 7.5)
        c.setFillColor(colors.HexColor("#7C3AED"))
        c.drawRightString(x + w - pad, y - 12 * mm, badge[:38])

    c.setStrokeColor(colors.HexColor("#E2E8F0"))
    c.line(x + pad, y - 16 * mm, x + w - pad, y - 16 * mm)

    c.setFont(PDF_FONT_BOLD, 8.8)
    c.setFillColor(colors.HexColor("#111827"))
    c.drawString(x + pad, y - 22 * mm, "Evolução")

    c.setFont(PDF_FONT, 9)
    c.setFillColor(colors.HexColor("#1E293B"))

    ly = y - 27 * mm
    for ln in _wrap_text(evolucao, PDF_FONT, 9, inner_w)[:13]:
        c.drawString(x + pad, ly, ln)
        ly -= 4.2 * mm

    return y - h - 5 * mm


def _new_pdf_page(pdf, timbre, page_w, page_h, titulo, subtitulo):
    _draw_timbre_page(pdf, timbre, page_w, page_h)
    box = _pdf_content_box(timbre, page_w, page_h)
    y = _draw_doc_title(pdf, box["x"], box["top"], box["w"], titulo, subtitulo)
    return box, y


def _gerar_pdf_evolucoes(
    *,
    conn,
    pacientes: Dict[str, Dict[str, Any]],
    timbre: dict,
    titulo: str,
    subtitulo: str,
    tipo_pdf: str,
    restrito: bool = False,
):
    buf = io.BytesIO()
    page_w, page_h = A4
    pdf = canvas.Canvas(buf, pagesize=A4)

    if not pacientes:
        box, y = _new_pdf_page(pdf, timbre, page_w, page_h, titulo, subtitulo)
        pdf.setFont(PDF_FONT_BOLD, 12)
        pdf.setFillColor(colors.HexColor("#334155"))
        pdf.drawString(box["x"], y, "Nenhum registro encontrado para os filtros selecionados.")
        pdf.showPage()
        pdf.save()
        buf.seek(0)
        return buf

    for pac in pacientes.values():
        box, y = _new_pdf_page(pdf, timbre, page_w, page_h, titulo, subtitulo)
        y = _draw_patient_box(pdf, box["x"], y, box["w"], pac, tipo_pdf)

        for item in pac["evos"]:
            badge = ""

            if restrito:
                if item.get("visibilidade") == "somente_eu":
                    badge = "Restrita • Somente eu"
                elif item.get("visibilidade") in ("cbos", "cbo"):
                    badge = "Restrita • CBO autorizado"
                else:
                    badge = "Restrita"

            needed_h = _registro_height(item.get("evo", "—"), box["w"]) + 8 * mm

            if y - needed_h < box["bottom"]:
                pdf.showPage()
                box, y = _new_pdf_page(pdf, timbre, page_w, page_h, titulo, subtitulo)
                y = _draw_patient_box(pdf, box["x"], y, box["w"], pac, tipo_pdf)

            y = _draw_registro_pdf(
                pdf,
                box["x"],
                y,
                box["w"],
                profissional=item.get("prof", "—"),
                cbo=item.get("cbo", ""),
                cbo_desc=_cbo_desc(conn, item.get("cbo", "")),
                data_atendimento=item.get("data", "—"),
                evolucao=item.get("evo", "—"),
                badge=badge,
            )

        pdf.showPage()

    pdf.save()
    buf.seek(0)
    return buf



@registros_bp.get("/evolucoes/pdf")
@require_permission("registros", "exportar")
def exportar_evolucoes_pdf():

    clinica_id = _clinica_id_atual()

    filtros = {
        "q": (request.args.get("q") or "").strip(),
        "prof": (request.args.get("prof") or "").strip(),
        "data_ini": (request.args.get("data_ini") or "").strip(),
        "data_fim": (request.args.get("data_fim") or "").strip(),
        "status": (request.args.get("status") or "").strip(),
        "sexo": (request.args.get("sexo") or "").strip(),
        "cid": (request.args.get("cid") or "").strip(),
        "cidade": (request.args.get("cidade") or "").strip(),
    }

    paciente_id = (request.args.get("paciente_id") or "").strip()

    try:
        limit_evos = int(request.args.get("limit_evos", 5000))
    except Exception:
        limit_evos = 5000

    conn = conectar_db()

    try:
        ensure_registros_schema(conn)

        if not _has_table(conn, "atendimentos"):
            return jsonify({"ok": False, "erro": "Tabela atendimentos inexistente"}), 500

        cur = conn.cursor()

        sql, params = _montar_query_atendimentos(
            conn,
            q=filtros["q"],
            prof=filtros["prof"],
            data_ini=filtros["data_ini"],
            data_fim=filtros["data_fim"],
            status=filtros["status"],
            sexo=filtros["sexo"],
            cid=filtros["cid"],
            cidade=filtros["cidade"],
            limit=limit_evos,
            incluir_evolucao_oculta=False,
        )

        cur.execute(sql, params)
        rows = _rows_to_dicts(cur, cur.fetchall())

        pacientes: Dict[str, Dict[str, Any]] = {}

        for r in rows:

            pid = str(r.get("paciente_id") or "0")

            if paciente_id and pid != paciente_id:
                continue

            nome = (
                r.get("pac__nome")
                or r.get("paciente_nome")
                or r.get("nome")
                or "—"
            )

            cpf = (
                r.get("pac__cpf")
                or r.get("cpf")
                or ""
            )

            cns = (
                r.get("pac__cns")
                or r.get("cns")
                or ""
            )

            nasc = (
                r.get("pac__nascimento")
                or r.get("nascimento")
                or ""
            )

            if pid not in pacientes:
                pacientes[pid] = {
                    "nome": nome,
                    "cpf": _fmt_cpf(cpf),
                    "cns": _fmt_cns(cns),
                    "nasc": nasc,
                    "status": r.get("status") or "—",
                    "mod": r.get("mod") or "—",
                    "evos": [],
                }

            evo = _safe_str(r.get("evolucao") or "")

            if not evo:
                continue

            pacientes[pid]["evos"].append({
                "data": (
                    r.get("data_atendimento")
                    or r.get("data")
                    or r.get("created_at")
                    or ""
                ),
                "prof": (
                    r.get("profissional_nome")
                    or r.get("nome_profissional")
                    or r.get("profissional")
                    or "—"
                ),
                "cbo": (
                    r.get("cbo_profissional")
                    or r.get("profissional_cbo")
                    or ""
                ),
                "evo": evo,
            })

        timbre = _buscar_timbre_pdf(conn, clinica_id)

        buf = _gerar_pdf_evolucoes(
            conn=conn,
            pacientes=pacientes,
            timbre=timbre,
            titulo="Relatório de Evoluções Gerais",
            subtitulo=f"Gerado em {date.today().strftime('%d/%m/%Y')}",
            tipo_pdf="Evoluções gerais",
            restrito=False,
        )

        registrar_log(
            modulo="registros",
            acao="exportar",
            entidade="pdf",
            descricao="Exportou PDF de evoluções gerais.",
            detalhes={
                "clinica_id": clinica_id,
                "total_pacientes": len(pacientes),
            },
        )

        return send_file(
            buf,
            as_attachment=True,
            download_name=f"evolucoes_gerais_{date.today().isoformat()}.pdf",
            mimetype="application/pdf",
        )

    except Exception as e:
        log_erro(
            "registros",
            e,
            entidade="pdf",
            descricao="Erro ao exportar PDF geral.",
            detalhes={"clinica_id": clinica_id},
        )

        return jsonify({
            "ok": False,
            "erro": str(e),
        }), 500

    finally:
        conn.close()



@registros_bp.get("/evolucoes-restritas/pdf")
@require_permission("registros", "exportar")
def exportar_evolucoes_restritas_pdf():
    clinica_id = _clinica_id_atual()
    paciente_id = (request.args.get("paciente_id") or "").strip()

    filtros = {
        "q": (request.args.get("q") or "").strip(),
        "prof": (request.args.get("prof") or "").strip(),
        "data_ini": (request.args.get("data_ini") or "").strip(),
        "data_fim": (request.args.get("data_fim") or "").strip(),
        "status": (request.args.get("status") or "").strip(),
        "sexo": (request.args.get("sexo") or "").strip(),
        "cid": (request.args.get("cid") or "").strip(),
        "cidade": (request.args.get("cidade") or "").strip(),
    }

    try:
        limit_evos = int(request.args.get("limit_evos", 5000))
    except Exception:
        limit_evos = 5000

    conn = conectar_db()

    try:
        ensure_registros_schema(conn)

        if not _has_table(conn, "atendimentos"):
            return jsonify({"ok": False, "erro": "Tabela atendimentos inexistente."}), 500

        if not _has_table(conn, "atendimento_evolucoes_ocultas"):
            return jsonify({"ok": False, "erro": "Tabela atendimento_evolucoes_ocultas inexistente."}), 500

        user = _resolve_logged_user(conn)
        uid = user.get("id") or 0
        user_cbo = _only_digits(user.get("cbo") or "")

        a_cols = _get_columns(conn, "atendimentos")

        col_a_id = _pick_att_col(a_cols, "id")
        col_clinica = _pick_att_col(a_cols, "clinica_id")
        col_pid = _pick_att_col(a_cols, "paciente_id")
        col_nome = _pick_att_col(a_cols, "paciente_nome", "nome_paciente", "nome")
        col_data = _pick_att_col(a_cols, "data_atendimento", "data", "data_iso", "created_at")
        col_prof_nome = _pick_att_col(a_cols, "profissional_nome", "nome_profissional", "profissional")
        col_cbo = _pick_att_col(a_cols, "cbo_profissional", "profissional_cbo", "cbo")
        col_mod = _pick_att_col(a_cols, "mod", "modalidade")
        col_status = _pick_att_col(a_cols, "status", "situacao", "comparecimento")
        col_cpf = _pick_att_col(a_cols, "paciente_cpf", "cpf")
        col_cns = _pick_att_col(a_cols, "paciente_cns", "cns", "cartao_sus")
        col_nasc = _pick_att_col(a_cols, "paciente_nascimento", "nascimento", "data_nascimento")

        if not (col_a_id and col_pid and col_nome and col_data):
            return jsonify({
                "ok": False,
                "erro": "Tabela atendimentos incompleta para PDF de evoluções restritas."
            }), 500

        d_expr = _sql_date_expr(f"a.{col_data}")

        where = [
            "TRIM(COALESCE(eo.evolucao_oculta, '')) <> ''",
            "(eo.clinica_id IS NULL OR eo.clinica_id = %s)",
            """
            (
                (
                    eo.visibilidade = 'somente_eu'
                    AND eo.profissional_id = %s
                )
                OR
                (
                    eo.visibilidade IN ('cbos', 'cbo')
                    AND %s <> ''
                    AND POSITION(
                        ',' || %s || ','
                        IN ',' || REGEXP_REPLACE(COALESCE(eo.cbos_autorizados, ''), '\\s+', '', 'g') || ','
                    ) > 0
                )
            )
            """
        ]

        params = [clinica_id, uid, user_cbo, user_cbo]

        if col_clinica:
            where.append(f"a.{col_clinica} = %s")
            params.append(clinica_id)

        if paciente_id:
            where.append(f"a.{col_pid} = %s")
            params.append(int(paciente_id) if paciente_id.isdigit() else paciente_id)

        data_ini = _norm_date_param(filtros["data_ini"])
        data_fim = _norm_date_param(filtros["data_fim"])

        if data_ini and data_fim:
            where.append(f"{d_expr} BETWEEN %s::date AND %s::date")
            params.extend([data_ini, data_fim])
        elif data_ini:
            where.append(f"{d_expr} = %s::date")
            params.append(data_ini)
        elif data_fim:
            where.append(f"{d_expr} = %s::date")
            params.append(data_fim)

        if filtros["q"]:
            like = f"%{filtros['q']}%"
            busca = []

            for c in [col_nome, col_cpf, col_cns]:
                if c:
                    busca.append(f"a.{c}::text ILIKE %s")
                    params.append(like)

            busca.append("eo.evolucao_oculta ILIKE %s")
            params.append(like)

            where.append("(" + " OR ".join(busca) + ")")

        if filtros["prof"]:
            col_prof_id = _pick_att_col(a_cols, "profissional_id", "prof_id", "id_profissional")

            if filtros["prof"].isdigit() and col_prof_id:
                where.append(f"a.{col_prof_id} = %s")
                params.append(int(filtros["prof"]))
            elif col_prof_nome:
                where.append(f"a.{col_prof_nome} ILIKE %s")
                params.append(f"%{filtros['prof']}%")

        if filtros["status"] and col_status:
            where.append(f"a.{col_status} = %s")
            params.append(filtros["status"])

        if filtros["sexo"]:
            col_sexo = _pick_att_col(a_cols, "sexo", "sex")
            if col_sexo:
                where.append(f"a.{col_sexo} = %s")
                params.append(filtros["sexo"])

        if filtros["cid"]:
            col_cid = _pick_att_col(a_cols, "cid", "cid_principal", "cid10")
            if col_cid:
                where.append(f"a.{col_cid} ILIKE %s")
                params.append(f"%{filtros['cid']}%")

        if filtros["cidade"]:
            col_cidade = _pick_att_col(a_cols, "cidade", "municipio", "cidade_paciente", "municipio_paciente")
            if col_cidade:
                where.append(f"a.{col_cidade} ILIKE %s")
                params.append(f"%{filtros['cidade']}%")

        sql = f"""
            SELECT
                a.{col_pid} AS paciente_id,
                a.{col_nome} AS paciente_nome,
                a.{col_data} AS data_atendimento,
                eo.evolucao_oculta AS evolucao,
                eo.visibilidade AS visibilidade,
                COALESCE(eo.profissional_nome, {f"a.{col_prof_nome}" if col_prof_nome else "''"}, '') AS profissional_nome,
                COALESCE(eo.profissional_cbo, {f"a.{col_cbo}" if col_cbo else "''"}, '') AS profissional_cbo,
                {f"COALESCE(a.{col_mod}, '')" if col_mod else "''"} AS paciente_mod,
                {f"COALESCE(a.{col_status}, '')" if col_status else "''"} AS paciente_status,
                {f"COALESCE(a.{col_cpf}, '')" if col_cpf else "''"} AS paciente_cpf,
                {f"COALESCE(a.{col_cns}, '')" if col_cns else "''"} AS paciente_cns,
                {f"COALESCE(a.{col_nasc}, '')" if col_nasc else "''"} AS paciente_nascimento
            FROM atendimento_evolucoes_ocultas eo
            INNER JOIN atendimentos a ON a.{col_a_id} = eo.atendimento_id
            WHERE {" AND ".join(where)}
            ORDER BY a.{col_pid} ASC, {d_expr} DESC, eo.id DESC
            LIMIT %s
        """

        cur = conn.cursor()
        cur.execute(sql, params + [limit_evos])
        rows = _rows_to_dicts(cur, cur.fetchall())

        pacientes: Dict[str, Dict[str, Any]] = {}

        for r in rows:
            pid = str(r.get("paciente_id") or "0")

            if pid not in pacientes:
                pacientes[pid] = {
                    "nome": _safe_str(r.get("paciente_nome")) or "—",
                    "cpf": _fmt_cpf(r.get("paciente_cpf", "")),
                    "cns": _fmt_cns(r.get("paciente_cns", "")),
                    "nasc": _safe_str(r.get("paciente_nascimento", "")),
                    "status": _safe_str(r.get("paciente_status", "")) or "—",
                    "mod": _safe_str(r.get("paciente_mod", "")) or "—",
                    "evos": [],
                }

            pacientes[pid]["evos"].append({
                "data": _safe_str(r.get("data_atendimento", "")) or "—",
                "prof": _safe_str(r.get("profissional_nome", "")) or "—",
                "cbo": _safe_str(r.get("profissional_cbo", "")) or "",
                "evo": _safe_str(r.get("evolucao", "")) or "—",
                "visibilidade": _safe_str(r.get("visibilidade", "")),
            })

        timbre = _buscar_timbre_pdf(conn, clinica_id)

        buf = _gerar_pdf_evolucoes(
            conn=conn,
            pacientes=pacientes,
            timbre=timbre,
            titulo="Relatório de Evoluções Restritas",
            subtitulo=f"Gerado em {date.today().strftime('%d/%m/%Y')}",
            tipo_pdf="Evoluções restritas",
            restrito=True,
        )

        registrar_log(
            modulo="registros",
            acao="exportar",
            entidade="pdf",
            descricao="Exportou PDF de evoluções restritas.",
            detalhes={
                "clinica_id": clinica_id,
                "usuario_id": uid,
                "usuario_cbo": user_cbo,
                "total_pacientes": len(pacientes),
                "filtros": filtros,
            },
        )

        return send_file(
            buf,
            as_attachment=True,
            download_name=f"evolucoes_restritas_{date.today().isoformat()}.pdf",
            mimetype="application/pdf",
        )

    except Exception as e:
        log_erro(
            "registros",
            e,
            entidade="pdf",
            descricao="Erro ao exportar PDF de evoluções restritas.",
            detalhes={"clinica_id": clinica_id},
        )

        return jsonify({
            "ok": False,
            "erro": str(e),
        }), 500

    finally:
        conn.close()




# =============================================================================
# BPA-i XLSX
# =============================================================================

def _fmt_date_bpai_ddmmyyyy(v: str) -> str:
    s = _safe_str(v)
    if not s:
        return ""

    s10 = s[:10]

    if re.match(r"^\d{4}-\d{2}-\d{2}$", s10):
        y, m, d = s10.split("-")
        return f"{d}/{m}/{y}"

    if re.match(r"^\d{2}/\d{2}/\d{4}$", s10):
        return s10

    return s10


def _calc_idade_no_dia(nasc: str, data_at: str) -> str:
    n_iso = _norm_date_param(_safe_str(nasc)[:10])
    d_iso = _norm_date_param(_safe_str(data_at)[:10])

    try:
        ny, nm, nd = int(n_iso[:4]), int(n_iso[5:7]), int(n_iso[8:10])
        dy, dm, dd = int(d_iso[:4]), int(d_iso[5:7]), int(d_iso[8:10])
        return str(max(0, dy - ny - ((dm, dd) < (nm, nd))))
    except Exception:
        return ""


def _map_raca_to_codigo(raca: str) -> str:
    s = _safe_str(raca).lower()

    if not s:
        return "99"

    d = _only_digits(s)

    if d:
        return d.zfill(2)[:2]

    m = {
        "branca": "01", "branco": "01",
        "preta": "02", "preto": "02", "negra": "02", "negro": "02",
        "parda": "03", "pardo": "03",
        "amarela": "04", "amarelo": "04",
        "indigena": "05", "indígena": "05",
        "ignorado": "99", "ignorada": "99",
        "sem informacao": "99", "sem informação": "99",
        "nao informado": "99", "não informado": "99",
    }

    return m.get(s, "99")


BPAI_COLS = [
    "prd-ident", "prd-cnes", "prd-cnsmed", "prd-cbo", "prd-dtaten", "prd-pa",
    "prd-cnspac", "prd-sexo", "prd-ibge", "prd-cid", "prd-idade", "prd-qt",
    "prd-caten", "prd-naut", "prd-org", "prd-nmpac", "prd-dtnasc", "prd-raca",
    "prd-etnia", "prd-nac", "prd-srv", "prd-clf", "prd-equipe-seq", "prd-equipe-area",
    "prd-cnpj", "prd-cep-pcnte", "prd-lograd-pcnte", "prd-end-pcnte", "prd-compl-pcnte",
    "prd-num-pcnte", "prd-bairro-pcnte", "prd-ddtel-pcnte", "prd-email-pcnte", "prd-ine"
]


def _fetch_dict_by_id(conn, table: str, id_value: Any) -> dict:
    if not id_value or not _has_table(conn, table) or not _has_column(conn, table, "id"):
        return {}

    clinica_id = _clinica_id_atual()
    cols = _get_columns(conn, table)
    col_clinica = _pick_att_col(cols, "clinica_id")

    cur = conn.cursor()

    if col_clinica:
        cur.execute(f"SELECT * FROM {table} WHERE id = %s AND {col_clinica} = %s LIMIT 1", (id_value, clinica_id))
    else:
        cur.execute(f"SELECT * FROM {table} WHERE id = %s LIMIT 1", (id_value,))

    row = cur.fetchone()

    if not row:
        return {}

    names = [d[0] for d in cur.description]

    if isinstance(row, dict):
        return dict(row)

    return {names[i]: row[i] for i in range(min(len(names), len(row)))}


def _rows_bpai(conn, filtros: dict) -> list[dict]:
    sql, params = _montar_query_atendimentos(
        conn,
        q=filtros.get("q", ""),
        prof=filtros.get("prof", ""),
        data_ini=filtros.get("data_ini", ""),
        data_fim=filtros.get("data_fim", ""),
        status=filtros.get("status", ""),
        sexo=filtros.get("sexo", ""),
        cid=filtros.get("cid", ""),
        cidade=filtros.get("cidade", ""),
        limit=filtros.get("limit", 50000),
        incluir_evolucao_oculta=False,
    )

    cur = conn.cursor()
    cur.execute(sql, params)
    at_list = _rows_to_dicts(cur, cur.fetchall())

    out: list[dict] = []

    for a in at_list:
        paciente_id = a.get("paciente_id")
        prof_id = a.get("profissional_id")

        pac = _fetch_dict_by_id(conn, "pacientes", paciente_id) if paciente_id else {}
        prof = _fetch_dict_by_id(conn, "usuarios", prof_id) if prof_id else {}

        dt_at_raw = a.get("data_atendimento") or a.get("data") or a.get("created_at") or ""
        dt_aten = _fmt_date_bpai_ddmmyyyy(str(dt_at_raw))

        nasc_raw = pac.get("nascimento") or a.get("nascimento") or ""
        dt_nasc = _fmt_date_bpai_ddmmyyyy(str(nasc_raw))

        sexo_p = _safe_str(pac.get("sexo") or a.get("sexo") or "")
        cid_p = _safe_str(pac.get("cid") or a.get("cid") or "")
        idade = _safe_str(pac.get("idade") or "") or _calc_idade_no_dia(str(nasc_raw), str(dt_at_raw))

        cnspac = _only_digits(pac.get("cns") or "")
        nmpac = _safe_str(pac.get("nome") or a.get("nome") or "")
        cnsmed = _only_digits(prof.get("cns") or a.get("cns_profissional") or "")
        cbo = _only_digits(prof.get("cbo") or a.get("cbo_profissional") or "")

        codigo_sigtap = _safe_str(a.get("ap_codigo_sigtap") or a.get("codigo_sigtap") or "")

        row = {k: "" for k in BPAI_COLS}
        row["prd-ident"] = "03"
        row["prd-cnes"] = ""
        row["prd-cnsmed"] = cnsmed
        row["prd-cbo"] = cbo
        row["prd-dtaten"] = dt_aten
        row["prd-pa"] = codigo_sigtap
        row["prd-cnspac"] = cnspac
        row["prd-sexo"] = sexo_p[:1].upper() if sexo_p else ""
        row["prd-ibge"] = ""
        row["prd-cid"] = cid_p

        try:
            row["prd-idade"] = int(str(idade).strip())
        except Exception:
            row["prd-idade"] = ""

        row["prd-qt"] = "000001"
        row["prd-caten"] = "01"
        row["prd-org"] = "BPA"
        row["prd-nmpac"] = nmpac
        row["prd-dtnasc"] = dt_nasc
        row["prd-raca"] = _map_raca_to_codigo(pac.get("raca") or "")
        row["prd-nac"] = "010"
        row["prd-cep-pcnte"] = _only_digits(pac.get("cep") or "")
        row["prd-lograd-pcnte"] = "081"
        row["prd-end-pcnte"] = _safe_str(pac.get("logradouro") or pac.get("rua") or "")
        row["prd-compl-pcnte"] = _safe_str(pac.get("complemento") or "")
        row["prd-num-pcnte"] = _safe_str(pac.get("numero_casa") or pac.get("numero") or "")
        row["prd-bairro-pcnte"] = _safe_str(pac.get("bairro") or "")
        row["prd-ddtel-pcnte"] = _only_digits(pac.get("telefone1") or pac.get("telefone") or "")

        out.append(row)

    return out


@registros_bp.get("/exportar_bpai_xlsx")
@require_permission("registros", "exportar")
def exportar_bpai_xlsx():
    clinica_id = _clinica_id_atual()

    filtros = {
        "q": (request.args.get("q") or "").strip(),
        "prof": (request.args.get("prof") or "").strip(),
        "data_ini": (request.args.get("data_ini") or "").strip(),
        "data_fim": (request.args.get("data_fim") or "").strip(),
        "status": (request.args.get("status") or "").strip(),
        "sexo": (request.args.get("sexo") or "").strip(),
        "cid": (request.args.get("cid") or "").strip(),
        "cidade": (request.args.get("cidade") or "").strip(),
    }

    try:
        filtros["limit"] = int(request.args.get("limit", 50000))
    except ValueError:
        filtros["limit"] = 50000

    conn = conectar_db()

    try:
        try:
            conn.rollback()
        except Exception:
            pass

        ensure_registros_schema(conn)
        rows = _rows_bpai(conn, filtros)

    except Exception as e:
        log_erro(
            "registros",
            e,
            entidade="bpai",
            descricao="Erro ao exportar BPA-I XLSX.",
            detalhes={"clinica_id": clinica_id, "filtros": filtros},
        )
        return jsonify({"ok": False, "erro": str(e)}), 500

    finally:
        conn.close()

    try:
        from openpyxl import Workbook
    except ImportError:
        return jsonify({"error": "Instale o pacote 'openpyxl' para exportar XLSX."}), 500

    wb = Workbook()
    ws = wb.active
    ws.title = "BPAi"

    ws.append(BPAI_COLS)

    if not rows:
        ws.append([""] * len(BPAI_COLS))
    else:
        for r in rows:
            ws.append([r.get(c, "") for c in BPAI_COLS])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    registrar_log(
        modulo="registros",
        acao="exportar",
        entidade="bpai",
        descricao="Exportou BPA-I XLSX.",
        detalhes={
            "clinica_id": clinica_id,
            "total": len(rows),
            "filtros": filtros,
        },
    )

    filename = f"bpai_{date.today().isoformat()}.xlsx"

    return send_file(
        buf,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )