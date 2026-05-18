from __future__ import annotations

from datetime import datetime, timedelta
import io
import csv
from typing import Optional, Any

from flask import render_template, request, jsonify, url_for, send_file, session, abort

from . import agenda_bp
from db import conectar_db

try:
    from admin.modulos import require_permission
except Exception:
    def require_permission(modulo_codigo: str, acao: str = "ver"):
        def deco(fn):
            return fn
        return deco

try:
    from log import registrar_log, log_erro
except Exception:
    def registrar_log(*args, **kwargs): pass
    def log_erro(*args, **kwargs): pass


DOW_LABELS_PT = ["Domingo", "Segunda", "Terça", "Quarta", "Quinta", "Sexta", "Sábado"]
DOW_LABELS = [(str(i), nome) for i, nome in enumerate(DOW_LABELS_PT)]


# ============================================================
# CONTEXTO
# ============================================================

def _clinica_id_atual() -> int:
    clinica_id = session.get("clinica_id")
    if not clinica_id:
        abort(403)
    try:
        return int(clinica_id)
    except Exception:
        abort(403)


def _usuario_id_atual():
    return session.get("usuario_id") or session.get("user_id") or session.get("id")


# ============================================================
# CONEXÃO / HELPERS
# ============================================================

def _conn():
    conn = conectar_db()
    try:
        from psycopg.rows import dict_row
        conn.row_factory = dict_row
    except Exception:
        pass
    return conn


def _fetchone_dict(cur):
    row = cur.fetchone()
    if row is None:
        return None
    if isinstance(row, dict):
        return dict(row)
    try:
        return dict(row)
    except Exception:
        cols = [c[0] for c in cur.description] if cur.description else []
        return dict(zip(cols, row))


def _fetchall_dicts(cur):
    rows = cur.fetchall() or []
    out = []
    cols = [c[0] for c in cur.description] if cur.description else []
    for row in rows:
        if isinstance(row, dict):
            out.append(dict(row))
        else:
            try:
                out.append(dict(row))
            except Exception:
                out.append(dict(zip(cols, row)))
    return out


def _has_table(conn, table: str) -> bool:
    cur = conn.cursor()
    cur.execute("""
        SELECT 1
          FROM information_schema.tables
         WHERE table_schema = 'public'
           AND table_name = %s
         LIMIT 1;
    """, (table,))
    return cur.fetchone() is not None


def _has_column(conn, table: str, column: str) -> bool:
    cur = conn.cursor()
    cur.execute("""
        SELECT 1
          FROM information_schema.columns
         WHERE table_schema = 'public'
           AND table_name = %s
           AND column_name = %s
         LIMIT 1;
    """, (table, column))
    return cur.fetchone() is not None


def _ensure_col(conn, table: str, col: str, ddl: str):
    cur = conn.cursor()
    cur.execute(f"ALTER TABLE {table} ADD COLUMN IF NOT EXISTS {col} {ddl};")
    conn.commit()


def _ensure_agendamentos_table(conn):
    cur = conn.cursor()
    cur.execute("""
        CREATE TABLE IF NOT EXISTS agendamentos (
            id SERIAL PRIMARY KEY,
            clinica_id INTEGER,

            paciente_id INTEGER,
            paciente TEXT,
            profissional TEXT,
            profissional_cpf TEXT,

            inicio TIMESTAMP,
            fim TIMESTAMP,

            dia TEXT,
            observacao TEXT,

            status TEXT DEFAULT 'ativo',

            recorrente INTEGER DEFAULT 0,
            serie_uid TEXT,
            dow_dom INTEGER,

            valor_sessao NUMERIC(12,2),

            criado_em TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            atualizado_em TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        );
    """)
    conn.commit()


def ensure_schema_agenda(conn):
    _ensure_agendamentos_table(conn)

    _ensure_col(conn, "agendamentos", "clinica_id", "INTEGER")
    _ensure_col(conn, "agendamentos", "paciente_id", "INTEGER")
    _ensure_col(conn, "agendamentos", "valor_sessao", "NUMERIC(12,2)")
    _ensure_col(conn, "agendamentos", "dia", "TEXT")
    _ensure_col(conn, "agendamentos", "criado_em", "TIMESTAMP DEFAULT CURRENT_TIMESTAMP")
    _ensure_col(conn, "agendamentos", "atualizado_em", "TIMESTAMP DEFAULT CURRENT_TIMESTAMP")

    cur = conn.cursor()

    cur.execute("CREATE INDEX IF NOT EXISTS idx_agendamentos_clinica ON agendamentos(clinica_id);")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_agendamentos_clinica_inicio ON agendamentos(clinica_id, inicio);")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_agendamentos_clinica_prof ON agendamentos(clinica_id, profissional_cpf);")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_agendamentos_clinica_paciente ON agendamentos(clinica_id, paciente_id);")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_agendamentos_clinica_dow_hora ON agendamentos(clinica_id, dow_dom, inicio);")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_agendamentos_status ON agendamentos(status);")

    conn.commit()


def _parse_hhmm(s: str) -> Optional[tuple[int, int]]:
    try:
        hh, mm = s.strip().split(":")
        hh, mm = int(hh), int(mm)
        if 0 <= hh <= 23 and 0 <= mm <= 59:
            return hh, mm
    except Exception:
        pass
    return None


def _next_date_for_dow(ref: datetime, dow_dom: int) -> datetime:
    today_w = (ref.weekday() + 1) % 7
    delta = (dow_dom - today_w) % 7
    return (ref + timedelta(days=delta)).replace(hour=0, minute=0, second=0, microsecond=0)


def _combine_date_time(day: datetime, hh: int, mm: int) -> datetime:
    return day.replace(hour=hh, minute=mm, second=0, microsecond=0)


def _to_dt(val):
    if val is None:
        return None
    if isinstance(val, datetime):
        return val
    try:
        return datetime.fromisoformat(str(val))
    except Exception:
        return None


def _usuario_by_cpf(conn, cpf: str):
    clinica_id = _clinica_id_atual()

    if not _has_table(conn, "usuarios"):
        return None

    cur = conn.cursor()
    cur.execute("""
        SELECT nome, cpf
          FROM usuarios
         WHERE TRIM(COALESCE(cpf, '')) = TRIM(%s)
           AND COALESCE(is_active, TRUE) = TRUE
           AND COALESCE(clinica_id, %s) = %s
         LIMIT 1;
    """, (cpf, clinica_id, clinica_id))

    r = _fetchone_dict(cur)
    if not r:
        return None

    return {"nome": r["nome"], "cpf": r["cpf"]}


# ============================================================
# PÁGINA
# ============================================================

@agenda_bp.route("/", methods=["GET"], endpoint="visualizar_agenda")
@require_permission("agenda", "ver")
def agenda_form():
    clinica_id = _clinica_id_atual()
    pacientes = []

    try:
        with _conn() as conn:
            ensure_schema_agenda(conn)

            cur = conn.cursor()
            try:
                cur.execute("""
                    SELECT id, nome
                      FROM pacientes
                     WHERE COALESCE(clinica_id, %s) = %s
                     ORDER BY nome;
                """, (clinica_id, clinica_id))
                pacientes = [{"id": r["id"], "nome": r["nome"]} for r in _fetchall_dicts(cur)]
            except Exception:
                pacientes = []

        registrar_log(
            modulo="agenda",
            acao="visualizar",
            entidade="agenda",
            descricao="Abriu a agenda.",
            detalhes={"clinica_id": clinica_id, "total_pacientes_combo": len(pacientes)},
        )

        return render_template("agenda.html", pacientes=pacientes, dow_labels=DOW_LABELS)

    except Exception as e:
        log_erro(
            "agenda",
            e,
            entidade="agenda",
            descricao="Erro ao abrir agenda.",
            detalhes={"clinica_id": clinica_id},
        )
        return f"Erro ao abrir agenda: {e}", 500


# ============================================================
# SALVAR AGENDAMENTO
# ============================================================

@agenda_bp.route("/", methods=["POST"], endpoint="agenda_salvar")
@require_permission("agenda", "editar")
def agenda_salvar():
    clinica_id = _clinica_id_atual()
    data = request.get_json(silent=True) or {}

    pac_id = data.get("paciente_id")
    pac_nome = (data.get("paciente_nome") or "").strip()
    dia = (data.get("dia") or "").strip()
    hora_de = (data.get("hora_de") or "").strip()
    hora_ate = (data.get("hora_ate") or "").strip()
    prof_cpf = (data.get("profissional_cpf") or "").strip()
    observacao = (data.get("observacao") or "").strip()

    is_vago = pac_nome.upper() == "VAGO"

    if not pac_id and not pac_nome and not is_vago:
        return jsonify({"error": "Informe o paciente."}), 400

    if not prof_cpf:
        return jsonify({"error": "Profissional é obrigatório."}), 400

    if not dia.isdigit() or not (0 <= int(dia) <= 6):
        return jsonify({"error": "Dia inválido. Use 0..6 (0=Dom..6=Sáb)."}), 400

    hhmm_de = _parse_hhmm(hora_de)
    if not hhmm_de:
        return jsonify({"error": "Horário inicial inválido. Use HH:MM."}), 400

    hhmm_ate = _parse_hhmm(hora_ate) if hora_ate else None

    try:
        with _conn() as conn:
            ensure_schema_agenda(conn)
            cur = conn.cursor()

            pid = None
            nome = None

            if is_vago:
                pid = None
                nome = "VAGO"
            else:
                if pac_id:
                    cur.execute("""
                        SELECT id, nome
                          FROM pacientes
                         WHERE id = %s
                           AND COALESCE(clinica_id, %s) = %s
                         LIMIT 1;
                    """, (pac_id, clinica_id, clinica_id))
                    r = _fetchone_dict(cur)
                    if not r:
                        return jsonify({"error": "Paciente não encontrado nesta clínica."}), 404
                    pid, nome = int(r["id"]), r["nome"]
                else:
                    cur.execute("""
                        SELECT id, nome
                          FROM pacientes
                         WHERE TRIM(UPPER(nome)) = TRIM(UPPER(%s))
                           AND COALESCE(clinica_id, %s) = %s
                         LIMIT 1;
                    """, (pac_nome, clinica_id, clinica_id))
                    r = _fetchone_dict(cur)
                    if not r:
                        return jsonify({"error": "Paciente não encontrado nesta clínica."}), 404
                    pid, nome = int(r["id"]), r["nome"]

            u = _usuario_by_cpf(conn, prof_cpf)
            if not u:
                return jsonify({"error": "Profissional não encontrado, inativo ou fora desta clínica."}), 400

            prof_nome = u["nome"]
            prof_cpf = u["cpf"]

            dow = int(dia)
            today = datetime.now()
            base_day = _next_date_for_dow(today, dow)

            if hhmm_ate:
                dur_min = (hhmm_ate[0] - hhmm_de[0]) * 60 + (hhmm_ate[1] - hhmm_de[1])
                if dur_min <= 0:
                    return jsonify({"error": "Horário final deve ser após o inicial."}), 400
            else:
                dur_min = 30

            ini_dt = _combine_date_time(base_day, hhmm_de[0], hhmm_de[1])
            fim_dt = ini_dt + timedelta(minutes=dur_min)
            dow_dom_val = ((ini_dt.weekday() + 1) % 7)

            try:
                dia_label = DOW_LABELS_PT[dow_dom_val]
            except Exception:
                dia_label = ""

            cur.execute("""
                INSERT INTO agendamentos (
                    clinica_id,
                    paciente_id,
                    paciente,
                    profissional,
                    profissional_cpf,
                    inicio,
                    fim,
                    dia,
                    observacao,
                    status,
                    recorrente,
                    serie_uid,
                    dow_dom,
                    valor_sessao,
                    criado_em,
                    atualizado_em
                )
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, 'ativo', 0, NULL, %s, NULL, CURRENT_TIMESTAMP, CURRENT_TIMESTAMP)
                RETURNING id;
            """, (
                clinica_id,
                pid,
                nome,
                prof_nome,
                prof_cpf,
                ini_dt,
                fim_dt,
                dia_label,
                observacao or None,
                dow_dom_val,
            ))

            created = cur.fetchone()
            created_id = int(created[0] if not isinstance(created, dict) else created["id"])
            conn.commit()

            redirect_url = None
            if pid is not None:
                try:
                    redirect_url = url_for("pacientes.visualizar_paciente", id=pid)
                except Exception:
                    redirect_url = None

        registrar_log(
            modulo="agenda",
            acao="criar",
            entidade="agendamentos",
            entidade_id=created_id,
            descricao="Criou agendamento.",
            detalhes={
                "clinica_id": clinica_id,
                "paciente_id": pid,
                "paciente": nome,
                "profissional": prof_nome,
                "profissional_cpf": prof_cpf,
                "inicio": ini_dt.isoformat(),
                "fim": fim_dt.isoformat(),
            },
        )

        return jsonify({
            "ok": True,
            "criado": created_id,
            "paciente": {"id": pid, "nome": nome},
            "redirect": redirect_url,
        }), 201

    except Exception as e:
        log_erro(
            "agenda",
            e,
            entidade="agendamentos",
            descricao="Erro ao criar agendamento.",
            detalhes={"clinica_id": clinica_id, "payload": data},
        )
        return jsonify({"error": "Falha ao criar agendamento", "detail": str(e)}), 500


# ============================================================
# API - PROFISSIONAIS
# ============================================================

@agenda_bp.get("/api/profissionais", endpoint="api_profissionais")
@require_permission("agenda", "ver")
def api_profissionais():
    clinica_id = _clinica_id_atual()

    try:
        with _conn() as conn:
            ensure_schema_agenda(conn)

            cur = conn.cursor()
            cur.execute("""
                SELECT nome, cpf
                  FROM usuarios
                 WHERE TRIM(COALESCE(cpf, '')) <> ''
                   AND COALESCE(is_active, TRUE) = TRUE
                   AND COALESCE(clinica_id, %s) = %s
                 ORDER BY nome;
            """, (clinica_id, clinica_id))
            rows = _fetchall_dicts(cur)

            out = [{"nome": r["nome"], "cpf": r["cpf"]} for r in rows]

        return jsonify(out)

    except Exception as e:
        log_erro(
            "agenda",
            e,
            entidade="usuarios",
            descricao="Erro ao listar profissionais para agenda.",
            detalhes={"clinica_id": clinica_id},
        )
        return jsonify({"error": "Falha ao listar profissionais", "detail": str(e)}), 500


# ============================================================
# API - AGREGADOS
# ============================================================

@agenda_bp.get("/api/agregados", endpoint="api_agregados")
@require_permission("agenda", "ver")
def api_agregados():
    clinica_id = _clinica_id_atual()

    try:
        qp = request.args

        prof_param = (qp.get("profissional_cpf") or qp.get("profissional") or "").strip()
        dia_param = (qp.get("dia") or qp.get("dia_semana") or "").strip()
        hora_de = (qp.get("hora_de") or qp.get("hora_ini") or "").strip()
        hora_ate = (qp.get("hora_ate") or qp.get("hora_fim") or "").strip()
        paciente_q = (qp.get("paciente") or qp.get("paciente_nome") or "").strip()
        idade_min = (qp.get("idade_min") or qp.get("idadeDe") or "").strip()
        idade_max = (qp.get("idade_max") or qp.get("idadeAte") or "").strip()
        cid_q = (qp.get("cid") or qp.get("cid10") or "").strip()

        with _conn() as conn:
            ensure_schema_agenda(conn)
            cur = conn.cursor()

            cur.execute("""
                SELECT column_name
                  FROM information_schema.columns
                 WHERE table_schema = 'public'
                   AND table_name = 'agendamentos';
            """)
            cols_ag = {r["column_name"] for r in _fetchall_dicts(cur)}

            has_pacientes = _has_table(conn, "pacientes")

            cols_pac = set()
            if has_pacientes:
                cur.execute("""
                    SELECT column_name
                      FROM information_schema.columns
                     WHERE table_schema = 'public'
                       AND table_name = 'pacientes';
                """)
                cols_pac = {r["column_name"] for r in _fetchall_dicts(cur)}

            has_financeiro = _has_table(conn, "financeiro_paciente_planos")
            has_dow = "dow_dom" in cols_ag
            has_prof_cpf = "profissional_cpf" in cols_ag
            has_pront = "prontuario" in cols_pac
            has_dia = "dia" in cols_ag
            has_ag_clinica = "clinica_id" in cols_ag
            has_pac_clinica = "clinica_id" in cols_pac

            nascimento_col = None
            if "nascimento" in cols_pac:
                nascimento_col = "p.nascimento"
            elif "data_nascimento" in cols_pac:
                nascimento_col = "p.data_nascimento"

            has_cid = "cid" in cols_pac

            dow_expr = "a.dow_dom" if has_dow else "EXTRACT(DOW FROM a.inicio)::integer"

            where_clauses = [
                "TRIM(COALESCE(a.paciente,'')) <> ''",
                "TRIM(COALESCE(a.profissional,'')) <> ''",
            ]
            params: list[Any] = []

            if has_ag_clinica:
                where_clauses.append("a.clinica_id = %s")
                params.append(clinica_id)

            if prof_param:
                if has_prof_cpf:
                    where_clauses.append("TRIM(COALESCE(a.profissional_cpf,'')) = TRIM(%s)")
                    params.append(prof_param)
                else:
                    where_clauses.append("TRIM(UPPER(a.profissional)) = TRIM(UPPER(%s))")
                    params.append(prof_param)

            if dia_param and dia_param.isdigit():
                where_clauses.append(f"{dow_expr} = %s")
                params.append(int(dia_param))

            if hora_de:
                where_clauses.append("TO_CHAR(a.inicio, 'HH24:MI') >= %s")
                params.append(hora_de)

            if hora_ate:
                where_clauses.append("TO_CHAR(a.inicio, 'HH24:MI') <= %s")
                params.append(hora_ate)

            if paciente_q:
                where_clauses.append("UPPER(a.paciente) LIKE '%%' || UPPER(%s) || '%%'")
                params.append(paciente_q)

            if nascimento_col and (idade_min or idade_max):
                age_expr = f"""
                CASE
                    WHEN {nascimento_col} IS NULL OR TRIM({nascimento_col}::text) = '' THEN NULL
                    WHEN {nascimento_col}::text ~ '^\\d{{4}}-\\d{{2}}-\\d{{2}}' THEN EXTRACT(YEAR FROM AGE(CURRENT_DATE, {nascimento_col}::date))::integer
                    WHEN {nascimento_col}::text ~ '^\\d{{2}}/\\d{{2}}/\\d{{4}}' THEN EXTRACT(YEAR FROM AGE(CURRENT_DATE, TO_DATE(SUBSTRING({nascimento_col}::text FROM 1 FOR 10), 'DD/MM/YYYY')))::integer
                    ELSE NULL
                END
                """
                if idade_min:
                    where_clauses.append(f"{age_expr} >= %s")
                    params.append(int(idade_min))
                if idade_max:
                    where_clauses.append(f"{age_expr} <= %s")
                    params.append(int(idade_max))

            if has_cid and cid_q:
                where_clauses.append("UPPER(COALESCE(p.cid,'')) LIKE '%%' || UPPER(%s) || '%%'")
                params.append(cid_q)

            where_sql = " AND ".join(where_clauses)

            select_pront = ", MIN(p.prontuario) AS prontuario" if has_pront else ""
            select_dia = ", MIN(a.dia) AS dia_data" if has_dia else ""

            if has_financeiro and has_pacientes:
                fp_has_clinica = _has_column(conn, "financeiro_paciente_planos", "clinica_id")
                join_fin_clinica = "AND COALESCE(fp.clinica_id, %s) = %s" if fp_has_clinica else ""
                if fp_has_clinica:
                    params_fin = [clinica_id, clinica_id]
                else:
                    params_fin = []

                select_fin = """
                    , MAX(COALESCE(fp.sessoes_contratadas, 0)) AS qtd_sessoes
                    , MAX(COALESCE(fp.sessoes_usadas, 0)) AS sessoes_usadas
                    , MAX(COALESCE(fp.combo_nome, fp.nome_plano, '')) AS nome_combo
                    , MAX(COALESCE(fp.status, '')) AS status_comercial
                """
                join_fin = f"""
                    LEFT JOIN financeiro_paciente_planos fp
                           ON fp.paciente_id = p.id
                          AND COALESCE(fp.status, 'ativo') = 'ativo'
                          {join_fin_clinica}
                """
            else:
                params_fin = []
                select_fin = """
                    , 0 AS qtd_sessoes
                    , 0 AS sessoes_usadas
                    , '' AS nome_combo
                    , '' AS status_comercial
                """
                join_fin = ""

            join_pacientes = ""
            if has_pacientes:
                join_pacientes = """
                    LEFT JOIN pacientes p
                           ON TRIM(UPPER(a.paciente)) = TRIM(UPPER(p.nome))
                """
                if has_pac_clinica:
                    join_pacientes += " AND p.clinica_id = a.clinica_id "

            sql = f"""
                SELECT
                    {dow_expr} AS dow,
                    TO_CHAR(a.inicio, 'HH24:MI') AS hora_ini,
                    a.profissional,
                    a.paciente,
                    COUNT(*) AS qtd_agendamentos,
                    MIN(a.id) AS any_id,
                    {("MIN(p.id) AS paciente_id, MIN(p.telefone) AS telefone" if has_pacientes else "NULL::integer AS paciente_id, NULL::text AS telefone")}
                    {select_pront if has_pacientes else ""}
                    {select_dia}
                    {select_fin}
                FROM agendamentos a
                {join_pacientes}
                {join_fin if has_pacientes else ""}
                WHERE {where_sql}
                GROUP BY {dow_expr}, TO_CHAR(a.inicio, 'HH24:MI'), a.profissional, a.paciente
                ORDER BY dow ASC, hora_ini ASC, a.profissional ASC, a.paciente ASC;
            """

            cur.execute(sql, params_fin + params)
            rows = _fetchall_dicts(cur)

            out = []
            for r in rows:
                dow = r["dow"]

                try:
                    dia_label = DOW_LABELS_PT[int(dow)]
                except Exception:
                    dia_label = "—"

                qtd_sessoes = int(r.get("qtd_sessoes") or 0)
                sessoes_usadas = int(r.get("sessoes_usadas") or 0)
                sessoes_restantes = max(qtd_sessoes - sessoes_usadas, 0)

                out.append({
                    "dia": dia_label,
                    "dia_num": int(dow) if dow is not None else None,
                    "dia_label": dia_label,
                    "hora_ini": r.get("hora_ini") or "—",
                    "profissional": r.get("profissional") or "—",
                    "paciente": r.get("paciente") or "—",
                    "telefone": r.get("telefone"),
                    "paciente_id": int(r["paciente_id"]) if r.get("paciente_id") is not None else None,
                    "prontuario": r.get("prontuario") if has_pront else None,
                    "qtd_agendamentos": int(r.get("qtd_agendamentos") or 0),
                    "qtd": qtd_sessoes,
                    "qtd_sessoes": qtd_sessoes,
                    "sessoes_usadas": sessoes_usadas,
                    "sessoes_restantes": sessoes_restantes,
                    "combo_nome": r.get("nome_combo") or "",
                    "status_comercial": r.get("status_comercial") or "",
                    "any_id": int(r.get("any_id") or 0),
                    "dia_data": r.get("dia_data") if has_dia else None,
                })

        registrar_log(
            modulo="agenda",
            acao="visualizar",
            entidade="agendamentos",
            descricao="Consultou agregados da agenda.",
            detalhes={"clinica_id": clinica_id, "total": len(out), "filtros": dict(request.args)},
        )

        return jsonify(out)

    except Exception as e:
        log_erro(
            "agenda",
            e,
            entidade="agendamentos",
            descricao="Erro ao listar agregados da agenda.",
            detalhes={"clinica_id": clinica_id, "filtros": dict(request.args)},
        )
        return jsonify({"error": "Falha ao listar agregados", "detail": str(e)}), 500


# ============================================================
# EXPORT AGREGADOS
# ============================================================

@agenda_bp.get("/api/agregados/export", endpoint="api_agregados_export")
@require_permission("agenda", "exportar")
def api_agregados_export():
    clinica_id = _clinica_id_atual()

    resp = api_agregados()

    if isinstance(resp, tuple):
        data, status = resp
        if status != 200:
            return resp
        items = data.get_json(force=True)
    else:
        items = resp.get_json(force=True)

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font

        wb = Workbook()
        ws = wb.active
        ws.title = "Agendamentos"

        headers = ["Dia", "Horário", "Profissional", "Paciente", "Qtd. registros", "Sessões", "Usadas", "Restantes", "Combo"]
        ws.append(headers)

        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")

        for it in items:
            ws.append([
                it.get("dia_label", "—"),
                it.get("hora_ini", "—"),
                it.get("profissional", "—"),
                it.get("paciente", "—"),
                it.get("qtd_agendamentos", 0),
                it.get("qtd_sessoes", 0),
                it.get("sessoes_usadas", 0),
                it.get("sessoes_restantes", 0),
                it.get("combo_nome", ""),
            ])

        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = min(max_len + 2, 50)

        bio = io.BytesIO()
        wb.save(bio)
        bio.seek(0)

        registrar_log(
            modulo="agenda",
            acao="exportar",
            entidade="agendamentos",
            descricao="Exportou agregados da agenda XLSX.",
            detalhes={"clinica_id": clinica_id, "total": len(items), "filtros": dict(request.args)},
        )

        fname = f"agendamentos_agregados_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
        return send_file(
            bio,
            as_attachment=True,
            download_name=fname,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    except ImportError:
        sio = io.StringIO()
        writer = csv.writer(sio, delimiter=";")
        writer.writerow(["Dia", "Horário", "Profissional", "Paciente", "Qtd. registros", "Sessões", "Usadas", "Restantes", "Combo"])

        for it in items:
            writer.writerow([
                it.get("dia_label", "—"),
                it.get("hora_ini", "—"),
                it.get("profissional", "—"),
                it.get("paciente", "—"),
                it.get("qtd_agendamentos", 0),
                it.get("qtd_sessoes", 0),
                it.get("sessoes_usadas", 0),
                it.get("sessoes_restantes", 0),
                it.get("combo_nome", ""),
            ])

        registrar_log(
            modulo="agenda",
            acao="exportar",
            entidade="agendamentos",
            descricao="Exportou agregados da agenda CSV.",
            detalhes={"clinica_id": clinica_id, "total": len(items), "filtros": dict(request.args)},
        )

        bio = io.BytesIO(sio.getvalue().encode("utf-8-sig"))
        fname = f"agendamentos_agregados_{datetime.now().strftime('%Y%m%d_%H%M')}.csv"
        return send_file(bio, as_attachment=True, download_name=fname, mimetype="text/csv")


# ============================================================
# API - OBTER AGENDAMENTO
# ============================================================

@agenda_bp.get("/api/agendamentos/<int:ag_id>", endpoint="api_get_agendamento")
@require_permission("agenda", "ver")
def api_get_agendamento(ag_id: int):
    clinica_id = _clinica_id_atual()

    try:
        with _conn() as conn:
            ensure_schema_agenda(conn)
            cur = conn.cursor()

            cur.execute("""
                SELECT id,
                       paciente_id,
                       paciente,
                       profissional,
                       profissional_cpf,
                       inicio,
                       fim,
                       dow_dom,
                       valor_sessao
                  FROM agendamentos
                 WHERE id = %s
                   AND clinica_id = %s
                 LIMIT 1;
            """, (ag_id, clinica_id))

            row = _fetchone_dict(cur)

            if not row:
                return jsonify({"error": "Agendamento não encontrado nesta clínica."}), 404

            dt_ini = _to_dt(row["inicio"])
            dt_fim = _to_dt(row["fim"])

            hora_de = dt_ini.strftime("%H:%M") if dt_ini else None
            hora_ate = dt_fim.strftime("%H:%M") if dt_fim else None

            if row.get("dow_dom") is not None:
                dia = int(row["dow_dom"])
            elif dt_ini:
                dia = (dt_ini.weekday() + 1) % 7
            else:
                dia = None

            return jsonify({
                "id": int(row["id"]),
                "paciente_id": row.get("paciente_id"),
                "paciente": row["paciente"],
                "profissional": row["profissional"],
                "profissional_cpf": row["profissional_cpf"],
                "dia": dia,
                "hora_de": hora_de,
                "hora_ate": hora_ate,
                "valor_sessao": float(row["valor_sessao"]) if row.get("valor_sessao") is not None else None,
            })

    except Exception as e:
        log_erro(
            "agenda",
            e,
            entidade="agendamentos",
            entidade_id=ag_id,
            descricao="Erro ao carregar agendamento.",
            detalhes={"clinica_id": clinica_id},
        )
        return jsonify({"error": "Falha ao carregar agendamento", "detail": str(e)}), 500


# ============================================================
# API - EDITAR AGENDAMENTO
# ============================================================

@agenda_bp.put("/api/agendamentos/<int:ag_id>", endpoint="api_editar_agendamento")
@require_permission("agenda", "editar")
def api_editar_agendamento(ag_id: int):
    clinica_id = _clinica_id_atual()
    data = request.get_json(silent=True) or {}

    dia_raw = (data.get("dia") or "").strip()
    hora_de = (data.get("hora_de") or "").strip()
    hora_ate = (data.get("hora_ate") or "").strip()
    novo_prof_cpf = (data.get("profissional_cpf") or "").strip()
    valor_sessao = data.get("valor_sessao", None)

    try:
        with _conn() as conn:
            ensure_schema_agenda(conn)
            cur = conn.cursor()

            cur.execute("""
                SELECT id,
                       paciente,
                       profissional,
                       profissional_cpf,
                       inicio,
                       fim,
                       dow_dom,
                       valor_sessao
                  FROM agendamentos
                 WHERE id = %s
                   AND clinica_id = %s
                 LIMIT 1;
            """, (ag_id, clinica_id))

            row = _fetchone_dict(cur)

            if not row:
                return jsonify({"error": "Agendamento não encontrado nesta clínica."}), 404

            dt_ini = _to_dt(row["inicio"])
            dt_fim = _to_dt(row["fim"])

            if not dt_ini:
                return jsonify({"error": "Não foi possível interpretar o horário inicial atual."}), 500

            dur_min_original = int((dt_fim - dt_ini).total_seconds() // 60) if dt_fim else 30

            novo_dia_semana = None
            if dia_raw:
                if not dia_raw.isdigit() or not (0 <= int(dia_raw) <= 6):
                    return jsonify({"error": "Valor de 'dia' inválido (use 0..6)."}), 400
                novo_dia_semana = int(dia_raw)

            if hora_de:
                hhmm_de = _parse_hhmm(hora_de)
                if not hhmm_de:
                    return jsonify({"error": "Horário inicial inválido. Use HH:MM."}), 400
                dt_ini_new = dt_ini.replace(hour=hhmm_de[0], minute=hhmm_de[1], second=0, microsecond=0)
            else:
                dt_ini_new = dt_ini

            if hora_ate:
                hhmm_ate = _parse_hhmm(hora_ate)
                if not hhmm_ate:
                    return jsonify({"error": "Horário final inválido. Use HH:MM."}), 400
                dt_fim_new = dt_ini_new.replace(hour=hhmm_ate[0], minute=hhmm_ate[1], second=0, microsecond=0)
                if dt_fim_new <= dt_ini_new:
                    return jsonify({"error": "Horário final deve ser após o inicial."}), 400
            else:
                dt_fim_new = dt_ini_new + timedelta(minutes=dur_min_original)

            if novo_dia_semana is not None:
                dow_atual_dom = (dt_ini_new.weekday() + 1) % 7
                delta_d = (novo_dia_semana - dow_atual_dom) % 7
                dt_ini_new = dt_ini_new + timedelta(days=delta_d)
                dt_fim_new = dt_fim_new + timedelta(days=delta_d)

            dow_dom_novo = (dt_ini_new.weekday() + 1) % 7
            dia_label = DOW_LABELS_PT[dow_dom_novo] if 0 <= dow_dom_novo <= 6 else ""

            novo_prof_nome = row["profissional"]
            novo_prof_cpf_resolvido = row["profissional_cpf"]

            if novo_prof_cpf:
                u = _usuario_by_cpf(conn, novo_prof_cpf)
                if not u:
                    return jsonify({"error": "Profissional não encontrado, inativo ou fora desta clínica."}), 400
                novo_prof_nome = u["nome"]
                novo_prof_cpf_resolvido = u["cpf"]

            val_float = None
            if valor_sessao is not None:
                try:
                    val_float = float(valor_sessao)
                except Exception:
                    return jsonify({"error": "Valor da sessão inválido."}), 400

            set_parts = [
                "inicio = %s",
                "fim = %s",
                "dia = %s",
                "dow_dom = %s",
                "profissional = %s",
                "profissional_cpf = %s",
                "atualizado_em = CURRENT_TIMESTAMP",
            ]

            params_up = [
                dt_ini_new,
                dt_fim_new,
                dia_label,
                dow_dom_novo,
                novo_prof_nome,
                novo_prof_cpf_resolvido,
            ]

            if val_float is not None:
                set_parts.append("valor_sessao = %s")
                params_up.append(val_float)

            params_up.extend([ag_id, clinica_id])

            cur.execute(f"""
                UPDATE agendamentos
                   SET {', '.join(set_parts)}
                 WHERE id = %s
                   AND clinica_id = %s;
            """, params_up)

            conn.commit()

        registrar_log(
            modulo="agenda",
            acao="editar",
            entidade="agendamentos",
            entidade_id=ag_id,
            descricao="Editou agendamento.",
            detalhes={
                "clinica_id": clinica_id,
                "inicio": dt_ini_new.isoformat(),
                "fim": dt_fim_new.isoformat(),
                "profissional": novo_prof_nome,
            },
        )

        return jsonify({"ok": True, "updated_id": ag_id})

    except Exception as e:
        log_erro(
            "agenda",
            e,
            entidade="agendamentos",
            entidade_id=ag_id,
            descricao="Erro ao editar agendamento.",
            detalhes={"clinica_id": clinica_id, "payload": data},
        )
        return jsonify({"error": "Falha ao editar agendamento", "detail": str(e)}), 500


# ============================================================
# API - EXCLUIR INDIVIDUAL
# ============================================================

@agenda_bp.delete("/api/agendamentos/<int:ag_id>", endpoint="api_excluir_agendamento")
@require_permission("agenda", "editar")
def api_excluir_agendamento(ag_id: int):
    clinica_id = _clinica_id_atual()

    try:
        with _conn() as conn:
            ensure_schema_agenda(conn)
            cur = conn.cursor()

            cur.execute("""
                SELECT id, paciente, profissional, inicio
                  FROM agendamentos
                 WHERE id = %s
                   AND clinica_id = %s
                 LIMIT 1;
            """, (ag_id, clinica_id))
            ag = _fetchone_dict(cur)

            if not ag:
                return jsonify({"error": "Agendamento não encontrado nesta clínica."}), 404

            cur.execute("""
                DELETE FROM agendamentos
                 WHERE id = %s
                   AND clinica_id = %s;
            """, (ag_id, clinica_id))

            deleted = cur.rowcount or 0
            conn.commit()

        registrar_log(
            modulo="agenda",
            acao="excluir",
            entidade="agendamentos",
            entidade_id=ag_id,
            descricao="Excluiu agendamento individual.",
            detalhes={"clinica_id": clinica_id, "agendamento": ag},
        )

        return jsonify({"ok": True, "deleted_id": ag_id, "deleted": deleted})

    except Exception as e:
        log_erro(
            "agenda",
            e,
            entidade="agendamentos",
            entidade_id=ag_id,
            descricao="Erro ao excluir agendamento individual.",
            detalhes={"clinica_id": clinica_id},
        )
        return jsonify({"error": "Falha ao excluir agendamento", "detail": str(e)}), 500


# ============================================================
# API - EXCLUIR GRUPO
# ============================================================

@agenda_bp.delete("/api/agregados", endpoint="api_excluir_grupo")
@require_permission("agenda", "editar")
def api_excluir_grupo():
    clinica_id = _clinica_id_atual()

    try:
        dow = request.args.get("dow", "").strip()
        hora_ini = request.args.get("hora_ini", "").strip()
        prof = (request.args.get("profissional") or "").strip()
        pac = (request.args.get("paciente") or "").strip()

        if not (dow.isdigit() and 0 <= int(dow) <= 6):
            return jsonify({"error": "Parâmetro 'dow' inválido (0..6)."}), 400

        if len(hora_ini) != 5 or ":" not in hora_ini:
            return jsonify({"error": "Parâmetro 'hora_ini' deve ser HH:MM."}), 400

        if not prof or not pac:
            return jsonify({"error": "Parâmetros 'profissional' e 'paciente' são obrigatórios."}), 400

        with _conn() as conn:
            ensure_schema_agenda(conn)
            cur = conn.cursor()

            cur.execute("""
                DELETE FROM agendamentos
                 WHERE clinica_id = %s
                   AND dow_dom = %s
                   AND TRIM(UPPER(profissional)) = TRIM(UPPER(%s))
                   AND TRIM(UPPER(paciente)) = TRIM(UPPER(%s))
                   AND TO_CHAR(inicio, 'HH24:MI') = %s;
            """, (clinica_id, int(dow), prof, pac, hora_ini))

            deleted = cur.rowcount or 0
            conn.commit()

        registrar_log(
            modulo="agenda",
            acao="excluir_grupo",
            entidade="agendamentos",
            descricao="Excluiu grupo agregado de agendamentos.",
            detalhes={
                "clinica_id": clinica_id,
                "dow": dow,
                "hora_ini": hora_ini,
                "profissional": prof,
                "paciente": pac,
                "deleted": deleted,
            },
        )

        return jsonify({"ok": True, "deleted": deleted})

    except Exception as e:
        log_erro(
            "agenda",
            e,
            entidade="agendamentos",
            descricao="Erro ao excluir grupo agregado.",
            detalhes={"clinica_id": clinica_id, "args": dict(request.args)},
        )
        return jsonify({"error": "Falha ao excluir grupo", "detail": str(e)}), 500